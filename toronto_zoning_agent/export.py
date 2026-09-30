# Copyright © 2026 Joshua Seaton. All rights reserved. Evaluation only; see LICENSE.
"""
Zoning Agent Excel workbook: output/zoning/zoning_reports.xlsx.

Rebuilt on every run from the report JSONs in output/zoning/ (the same contract that
drives the HTML reports), so it never parses its own earlier state. One workbook per
report folder, read top-down from an executive summary to the City's own records:

  Summary        one row per lot: as-of-right envelope, the lot's last decision and
                 permit, neighbour precedent, signs of change, data check, links
  <lot>          one brief per lot: bottom line, key numbers, as of right with by-law
                 links, the lot's history, precedent by project type, signs of change
  Applications   every neighbour Committee of Adjustment file (Excel table, Lot first)
  Permits        every neighbour building-permit project
  Development Applications  rezonings, site plans, subdivisions and condominiums
  Sources        datasets, data check per lot, how the numbers are made, limits

Every record row has two links: Documents, the same route as the HTML record card (the
file's Application Information Centre page when the AIC still carries it, else the
Committee of Adjustment staff, the permit status search, a building records request or
Development Review), and Raw City data, the exact Open Data row the numbers came from. Counts on the Summary and briefs
are COUNTIFS formulas over the Applications table, so each number can be traced to its
rows; their values are also cached in the file so previews that don't calculate
(Quick Look, pandas) still show them. Links are native hyperlinks, not =HYPERLINK()
formulas, so they work in Excel, Numbers, LibreOffice and Google Sheets.

Rebuild without rerunning the agent:  python -m toronto_zoning_agent.export
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
import zipfile
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from toronto_zoning_agent import rules as zr
from toronto_zoning_agent.paths import OUTPUT_DIR

REPORT_DIR = OUTPUT_DIR / "zoning"
WORKBOOK_NAME = "zoning_reports.xlsx"
NOTICE = "Parcis · © 2026 Joshua Seaton. All rights reserved. Evaluation copy."
WATERMARK = "Parcis · Evaluation copy"
DEFAULT_RADIUS_M = 250
LICENCE_URL = "https://open.toronto.ca/open-data-licence/"

# ── design tokens (DESIGN.md) ─────────────────────────────────────────────────
FONT = "Arial"
# Parcis Phase I palette
SLATE_BLUE = "4D6B84"   # title bands, links
SLATE = "2A3A4A"        # table and section headers
GROUP = "E4EDF4"        # group bands
CREAM = "F7F2EA"        # subtitle and note bands
SAND = "C5B9AD"         # hairlines
STONE = "8A7E72"        # secondary text
DRIFTWOOD = "4A3F36"
INK = "111820"          # Midnight
INK_2 = "4A3F36"        # Driftwood: labels and detail text
MUTED = STONE
HAIR = SAND
SURFACE_2 = CREAM
ZONE = "F0D77A"
ZONE_INK = "2A2410"
LINK = SLATE_BLUE
WHITE = "FFFFFF"
WARN_WASH, WARN_INK = "FFF4D6", "6B4A00"
SIGNAL = {  # fill, ink
    "Strong": ("DCFCE7", "166534"),
    "Mixed": ("FEF3C7", "92400E"),
    "Weak": ("FEE2E2", "991B1B"),
    "Sparse": ("F1F3F4", "374151"),
    "Unchecked": (WARN_WASH, WARN_INK),
}
SIGNAL_TEXT = {
    "Strong": "Strong precedent", "Mixed": "Mixed precedent", "Weak": "Weak precedent",
    "Sparse": "Too few decisions", "Unchecked": "Couldn't be checked",
}
OUTCOME = {  # JSON bucket -> (label in the workbook, fill, ink)
    "Approved": ("Approved", "E6F4EA", "166534"),
    "Refused": ("Refused", "FDECEC", "991B1B"),
    "Pending": ("Awaiting hearing", WARN_WASH, WARN_INK),
    "Other": ("Not decided", "F1F3F4", INK_2),
}
PROJECT_TYPES = ["New house", "Addition or alteration", "Pool, deck or exterior",
                 "Legalize existing work", "Severance or consent", "Other"]
STATE_TEXT = {"found": "Records found", "none": "Checked, none", "unchecked":
              "Couldn't be checked", "skipped": "Skipped this run"}

F_INT = "#,##0"
F_PCT = "0%"
F_MONEY = "$#,##0"
F_DATE = "d mmm yyyy"
F_FSI = "0.00"
F_STOREYS = "General"  # 3 -> "3", 2.5 -> "2.5"

THIN = Side(style="thin", color=HAIR)
BOTTOM = Border(bottom=THIN)
SLATE_RULE = Border(bottom=Side(style="medium", color=SLATE))

# Record-sheet table columns (the COUNTIFS below use these letters).
APP_COLS = [  # (header, width, number format)
    ("Lot", 20, None), ("Dist. (m)", 9, F_INT), ("Address", 22, None), ("File", 14, None),
    ("Documents", 22, None), ("Application", 14, None), ("Project type", 21, None), ("Storeys", 8, F_STOREYS),
    ("Outcome", 16, None), ("Decision on file", 22, None), ("Filed", 12, F_DATE),
    ("Hearing", 12, F_DATE), ("Weeks to hearing", 10, F_INT), ("Appeal", 14, None),
    ("Built under permit", 16, None), ("What was asked", 64, None),
    ("Map", 8, None), ("Raw City data", 12, None),
]
PERMIT_COLS = [
    ("Lot", 20, None), ("Dist. (m)", 9, F_INT), ("Address", 22, None), ("Permit", 12, None),
    ("Documents", 22, None), ("Category", 20, None), ("Permit types", 26, None), ("Structure", 18, None),
    ("Status", 16, None), ("Applied", 12, F_DATE), ("Issued", 12, F_DATE),
    ("Completed", 12, F_DATE), ("Est. cost", 13, F_MONEY), ("Cost note", 18, None),
    ("Units", 7, F_INT), ("Cites variance", 16, None), ("Description", 64, None),
    ("Map", 8, None), ("Raw City data", 12, None),
]
REZ_COLS = [
    ("Lot", 20, None), ("Dist. (m)", 9, F_INT), ("Address", 22, None),
    ("Application", 22, None), ("Documents", 22, None), ("Type", 30, None), ("Status", 22, None),
    ("Submitted", 12, F_DATE), ("Description", 70, None), ("Map", 8, None),
    ("Raw City data", 12, None),
]
_COL = {h: get_column_letter(i) for i, (h, _, _) in enumerate(APP_COLS, start=1)}
_TABLE_TOP = 4  # header row of every record table


# ── small helpers ─────────────────────────────────────────────────────────────

# Body point size: 10 on record sheets, 11 on the Summary and lot briefs (set while
# those are written). Sizes up to 12 shift with it; titles and tile numbers don't.
_BODY_PT = 10


def _font(size=10, bold=False, color=INK, italic=False, underline=None) -> Font:
    if size <= 12:
        size += _BODY_PT - 10
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic, underline=underline)


def _lh(h: float) -> float:
    """A row height measured at 10 pt, scaled to the current body size."""
    return h * _BODY_PT / 10


class _body_pt:
    def __init__(self, pt: int):
        self.pt = pt

    def __enter__(self):
        global _BODY_PT
        self.prev, _BODY_PT = _BODY_PT, self.pt

    def __exit__(self, *exc):
        global _BODY_PT
        _BODY_PT = self.prev


DEV_SHEET = "Development Applications"
OZ_NAME = "Official Plan Amendment / Rezoning"
DOC_TEXT = {"aic": "Open on City site", "coa_staff": "Request from CoA staff",
            "research": "Research Request", "status": "Check permit status",
            "records": "Building records request", "devreview": "developmentreview@toronto.ca"}


def doc_route(docs) -> Optional[dict]:
    """The first (best) document route for a record (data.document_routes)."""
    return (docs or [None])[0]


def doc_text(docs) -> str:
    r = doc_route(docs)
    return DOC_TEXT.get(r.get("action"), r.get("label", "")) if r and r.get("url") else ""


def _doc(docs) -> tuple:
    r = doc_route(docs)
    return ("link", (r or {}).get("url"), doc_text(docs))


def devapp_is_sign(d: dict, run_date: str) -> bool:
    """Signs of change: open applications, or ones submitted in the last 3 years."""
    if d.get("open"):
        return True
    sub, run = _date(d.get("submitted")), _date(run_date)
    if not sub or not run:
        return False
    try:
        cutoff = run.replace(year=run.year - 3)
    except ValueError:
        cutoff = run.replace(year=run.year - 3, day=28)
    return sub >= cutoff


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def _date(iso) -> Optional[date]:
    try:
        return date.fromisoformat(str(iso)[:10]) if iso else None
    except ValueError:
        return None


def _num(v) -> Optional[float]:
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    return int(f) if f.is_integer() else f


def _fmt_date(iso) -> str:
    d = _date(iso)
    return f"{d.day} {d:%b %Y}" if d else ""


def _fmt_month(iso) -> str:
    d = _date(iso)
    return f"{d:%b %Y}" if d else ""


def _plural(n: int, one: str, many: str) -> str:
    return f"{n} {one if n == 1 else many}"


def _q(s: str) -> str:
    """Excel string literal for a formula."""
    return '"' + str(s).replace('"', '""') + '"'


def lot_label(data: dict) -> str:
    r = data.get("radius_m") or DEFAULT_RADIUS_M
    return data.get("address", "") + ("" if r == DEFAULT_RADIUS_M else f" ({r} m)")


def _sheet_name(label: str, used: set) -> str:
    base = re.sub(r"[:\\/?*\[\]']", " ", label).strip()[:31] or "Lot"
    name, i = base, 2
    while name.lower() in used:
        suffix = f" {i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name.lower())
    return name


def _ref(sheet: str, cell: str = "A1") -> str:
    return f"'{sheet}'!{cell}"


def _lines(text: str, width_chars: float) -> int:
    """Wrapped line count for text in a cell `width_chars` wide (Arial, current body size)."""
    if not text:
        return 1
    per = max(8, int(width_chars * 1.1 * 10 / _BODY_PT))
    return sum(max(1, math.ceil(len(p) / per)) for p in str(text).split("\n"))


class _Book:
    """Workbook plus the cached values for the formulas written into it."""

    def __init__(self):
        self.wb = Workbook()
        normal = self.wb._named_styles["Normal"]
        normal.font = Font(name=FONT, size=10, color=INK)
        self.cache: dict[str, dict[str, object]] = {}

    def formula(self, ws, cell: str, formula: str, value, fmt: Optional[str] = None,
                font: Optional[Font] = None, align: Optional[Alignment] = None):
        c = ws[cell]
        c.value = "=" + formula
        if fmt:
            c.number_format = fmt
        c.font = font or _font()
        if align:
            c.alignment = align
        self.cache.setdefault(ws.title, {})[cell] = value
        return c


def _link(cell, target: str, text: str, internal: bool = False, size: int = 10,
          bold: bool = False) -> None:
    cell.value = text
    if internal:
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=target, display=text)
    else:
        cell.hyperlink = target
    cell.font = _font(size=size, bold=bold, color=LINK, underline="single")


def _merge(ws, rng: str, value=None, font: Optional[Font] = None, fill=None,
           align: Optional[Alignment] = None, fmt: Optional[str] = None):
    ws.merge_cells(rng)
    c = ws[rng.split(":")[0]]
    if value is not None:
        c.value = value
    if font:
        c.font = font
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    if fmt:
        c.number_format = fmt
    return c


def _title_band(ws, last_col: str, title: str, subtitle: str) -> None:
    subtitle = f"{subtitle}  ·  {NOTICE}"
    _merge(ws, f"A1:{last_col}1", title, _font(16, True, WHITE), _fill(SLATE_BLUE),
           Alignment(vertical="center", indent=1))
    ws.row_dimensions[1].height = 34
    _merge(ws, f"A2:{last_col}2", subtitle, _font(10, color=INK_2), _fill(CREAM),
           Alignment(vertical="center", indent=1, wrap_text=True))
    ws.row_dimensions[2].height = _lh(34)


def _section(ws, row: int, last_col: str, text: str, note: str = "") -> int:
    _merge(ws, f"A{row}:{last_col}{row}", text.upper(), _font(10, True, SLATE),
           align=Alignment(vertical="bottom"))
    for col in range(1, _col_index(last_col) + 1):
        ws.cell(row=row, column=col).border = SLATE_RULE
    ws.row_dimensions[row].height = 24
    if note:
        row += 1
        _merge(ws, f"A{row}:{last_col}{row}", note, _font(9, color=MUTED, italic=True),
               align=Alignment(wrap_text=True, vertical="top"))
        width = sum(ws.column_dimensions[get_column_letter(i)].width or 9
                    for i in range(1, _col_index(last_col) + 1))
        ws.row_dimensions[row].height = max(_lh(15), _lh(12) * _lines(note, width) + 4)
    return row + 1


def _col_index(letter: str) -> int:
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch) - 64)
    return n


def _print_setup(ws, landscape: bool = True, title_rows: Optional[str] = None) -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.oddFooter.left.text = NOTICE
    ws.oddHeader.right.text = WATERMARK
    ws.oddHeader.right.size = 8
    ws.oddHeader.right.color = "8A7E72"
    ws.oddFooter.left.size = 8
    ws.oddFooter.center.text = "&A"
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    if title_rows:
        ws.print_title_rows = title_rows


# ── facts derived from one report (same rules as the HTML page) ───────────────

def _records(data: dict) -> list:
    return data.get("records") or []


def _state(data: dict, key: str) -> str:
    return (data.get("completeness") or {}).get(key, "")


def _unchecked(data: dict, key: str) -> bool:
    return _state(data, key) == "unchecked"


def _counts(data: dict) -> dict:
    recs = _records(data)
    out = {"apps": len(recs), "approved": 0, "refused": 0, "pending": 0, "other": 0,
           "nh3": 0, "sever": 0}
    for r in recs:
        o = r.get("outcome")
        out[{"Approved": "approved", "Refused": "refused", "Pending": "pending"}.get(o, "other")] += 1
        if o == "Approved" and r.get("type") == "New house" and r.get("storeys") == 3:
            out["nh3"] += 1
        if r.get("type") == "Severance or consent":
            out["sever"] += 1
    out["decided"] = out["approved"] + out["refused"]
    out["rate"] = out["approved"] / out["decided"] if out["decided"] else ""
    return out


def _type_counts(data: dict, ptype: str) -> dict:
    rs = [r for r in _records(data) if r.get("type") == ptype]
    c = {k: sum(1 for r in rs if r.get("outcome") == k) for k in ("Approved", "Refused", "Pending", "Other")}
    c["n"] = len(rs)
    decided = c["Approved"] + c["Refused"]
    c["rate"] = c["Approved"] / decided if decided else ""
    return c


def _storey_counts(data: dict) -> dict:
    nh = [r for r in _records(data) if r.get("type") == "New house" and r.get("outcome") == "Approved"]
    out = {s: sum(1 for r in nh if r.get("storeys") == s) for s in (3, 2.5, 2)}
    out["other"] = sum(1 for r in nh if r.get("storeys") not in (None, 3, 2.5, 2))
    out["unknown"] = sum(1 for r in nh if r.get("storeys") is None)
    out["all"] = len(nh)
    return out


def _timing(data: dict) -> tuple[Optional[int], str, str]:
    """(weeks, short basis, long note) for the typical time to hearing."""
    s = data.get("stats") or {}
    t = s.get("timing") or {}
    cutoff = _fmt_month(t.get("cutoff"))
    if s.get("recent_hearing_weeks") is not None:
        w = s["recent_hearing_weeks"]
        note = (f"Median of {t.get('recent_n')} decided files filed since {cutoff} "
                f"(filing to hearing).")
        if s.get("all_hearing_weeks") is not None:
            note += f" All years: {s['all_hearing_weeks']} weeks ({t.get('all_n')} files)."
        return w, f"median, last 3 yrs ({t.get('recent_n')} files)", note
    if s.get("all_hearing_weeks") is not None:
        return (s["all_hearing_weeks"], f"median, all years ({t.get('all_n')} files)",
                f"Median of {t.get('all_n')} decided files, all years; fewer than 3 filed "
                f"since {cutoff}.")
    return None, "too few decisions", "Fewer than 3 decided files, so no typical time."


def _last_decision(data: dict) -> str:
    if _unchecked(data, "subject_coa"):
        return "Couldn't be checked"
    coa = sorted((data.get("subject") or {}).get("coa") or [],
                 key=lambda c: c.get("hearing") or c.get("filed") or "")
    decided = [c for c in coa if c.get("outcome") in ("Approved", "Refused")]
    if decided:
        c = decided[-1]
        return f"{c.get('app') or 'Application'} {c['outcome'].lower()} {_fmt_date(c.get('hearing'))}"
    if coa:
        return _plural(len(coa), "application on file, none decided", "applications on file, none decided")
    return "No variance on record"


def _permit_status(data: dict) -> str:
    if _unchecked(data, "subject_permits"):
        return "Couldn't be checked"
    ps = (data.get("subject") or {}).get("permits") or []
    p = next((p for p in ps if p.get("category") == "New building"), ps[0] if ps else None)
    if not p:
        return "No building permit on record"
    kind = "New-building permit" if p.get("category") == "New building" else "Building permit"
    if p.get("issued"):
        return f"{kind} issued {_fmt_month(p['issued'])}, {p.get('status_text') or p.get('status', '')}"
    return f"{kind} applied for {_fmt_month(p.get('applied'))}, {p.get('status_text', '')}"


def _data_check(data: dict) -> tuple[str, bool]:
    unchecked = data.get("unchecked_sources") or []
    if unchecked:
        return f"{_plural(len(unchecked), 'source', 'sources')} couldn't be checked", True
    n = sum(1 for v in (data.get("completeness") or {}).values() if v in ("found", "none"))
    return f"All {n} sources checked", False


def _tok(data: dict, key: str):
    return next((t for t in (data.get("zone") or {}).get("tokens") or [] if t.get("key") == key), None)


# ── Summary ───────────────────────────────────────────────────────────────────

SUMMARY_COLS = [  # (group, header, width)
    ("Lot", "Address", 22), ("Lot", "Ward", 18),
    ("As of right", "Zone", 25), ("As of right", "Max height (m)", 9), ("As of right", "Max FSI", 8),
    ("This lot", "Last decision", 27), ("This lot", "Permit", 30),
    ("Neighbour precedent", "Radius (m)", 8), ("Neighbour precedent", "Applications", 11),
    ("Neighbour precedent", "Approved", 9), ("Neighbour precedent", "Decided", 9),
    ("Neighbour precedent", "Approval rate", 9), ("Neighbour precedent", "Precedent", 17),
    ("Neighbour precedent", "3-storey new houses approved", 11),
    ("Neighbour precedent", "Weeks to hearing", 10),
    ("Signs of change", "Awaiting hearing", 9), ("Signs of change", "Severance attempts", 10),
    ("Signs of change", "Rezonings / OPAs (OZ)", 11),
    ("Data", "Data check", 23),
    ("Open", "Report", 12), ("Open", "Map", 7), ("Open", "Street View", 10),
]


def _write_summary(book: _Book, ws, reports: list, briefs: dict) -> None:
    last = get_column_letter(len(SUMMARY_COLS))
    run_dates = sorted({d.get("generated", "") for d in reports if d.get("generated")})
    when = _fmt_date(run_dates[-1]) if len(run_dates) == 1 else \
        f"{_fmt_date(run_dates[0])} to {_fmt_date(run_dates[-1])}" if run_dates else ""
    _title_band(ws, last, "Parcis Zoning Reports",
                f"City of Toronto  ·  {_plural(len(reports), 'lot', 'lots')}  ·  as-of-right "
                f"envelope, the lot's history and Committee of Adjustment precedent nearby  ·  "
                f"runs {when}  ·  source: City of Toronto Open Data")
    ws.row_dimensions[3].height = 8
    # group band
    groups: list[tuple[str, int, int]] = []
    for i, (g, _, _) in enumerate(SUMMARY_COLS, start=1):
        if groups and groups[-1][0] == g:
            groups[-1] = (g, groups[-1][1], i)
        else:
            groups.append((g, i, i))
    for g, a, b in groups:
        rng = f"{get_column_letter(a)}4:{get_column_letter(b)}4"
        if a != b:
            ws.merge_cells(rng)
        c = ws.cell(row=4, column=a, value=g.upper())
        c.font = _font(8, True, INK_2)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(a, b + 1):
            cell = ws.cell(row=4, column=col)
            cell.fill = _fill(GROUP)
            cell.border = Border(left=THIN if col == a else None, top=THIN)
    ws.row_dimensions[4].height = 18
    for i, (_, h, w) in enumerate(SUMMARY_COLS, start=1):
        c = ws.cell(row=5, column=i, value=h)
        c.font = _font(9, True, WHITE)
        c.fill = _fill(SLATE)
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="left" if i <= 7 or h in ("Precedent", "Data check") else "center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[5].height = 42

    col = {h: get_column_letter(i) for i, (_, h, _) in enumerate(SUMMARY_COLS, start=1)}
    lot_rng = f"Applications!${_COL['Lot']}:${_COL['Lot']}"
    out_rng = f"Applications!${_COL['Outcome']}:${_COL['Outcome']}"
    type_rng = f"Applications!${_COL['Project type']}:${_COL['Project type']}"
    st_rng = f"Applications!${_COL['Storeys']}:${_COL['Storeys']}"
    center = Alignment(horizontal="center", vertical="center")
    for n, data in enumerate(reports):
        row = 6 + n
        r = str(row)
        label = lot_label(data)
        sheet = briefs[label]
        zone = data.get("zone") or {}
        env = data.get("envelope") or {}
        cnt = _counts(data)
        fsi = _tok(data, "fsi") or _tok(data, "total_fsi")
        _link(ws[f"A{r}"], _ref(sheet), label, internal=True, bold=True)
        ws[f"B{r}"] = (data.get("ward") or "").replace("-", "–") or "—"
        ws[f"C{r}"] = zone.get("raw") or ("Couldn't be checked" if _unchecked(data, "zone_map")
                                         else "Not on the zoning map")
        h = _num(env.get("height_m"))
        ws[f"D{r}"] = h if h is not None else "—"
        ws[f"E{r}"] = _num(fsi["value"]) if fsi else "—"
        ws[f"E{r}"].number_format = F_FSI
        ws[f"F{r}"] = _last_decision(data)
        ws[f"G{r}"] = _permit_status(data)
        ws[f"H{r}"] = data.get("radius_m")
        ws[f"H{r}"].number_format = F_INT
        crit = f"$A{r}"
        if _unchecked(data, "neighbour_coa"):
            for h_ in ("Applications", "Approved", "Decided", "Approval rate",
                       "3-storey new houses approved", "Weeks to hearing", "Awaiting hearing",
                       "Severance attempts"):
                c = ws[f"{col[h_]}{r}"]
                c.value = "n/a ⚠"
                c.font = _font(10, True, WARN_INK)
        else:
            book.formula(ws, f"{col['Applications']}{r}", f"COUNTIFS({lot_rng},{crit})",
                         cnt["apps"], F_INT)
            book.formula(ws, f"{col['Approved']}{r}",
                         f"COUNTIFS({lot_rng},{crit},{out_rng},\"Approved\")", cnt["approved"], F_INT)
            book.formula(ws, f"{col['Decided']}{r}",
                         f"COUNTIFS({lot_rng},{crit},{out_rng},\"Approved\")"
                         f"+COUNTIFS({lot_rng},{crit},{out_rng},\"Refused\")", cnt["decided"], F_INT)
            book.formula(ws, f"{col['Approval rate']}{r}",
                         f"IF({col['Decided']}{r}>0,{col['Approved']}{r}/{col['Decided']}{r},\"\")",
                         cnt["rate"], F_PCT, _font(10, True))
            book.formula(ws, f"{col['3-storey new houses approved']}{r}",
                         f"COUNTIFS({lot_rng},{crit},{type_rng},\"New house\",{out_rng},\"Approved\",{st_rng},3)",
                         cnt["nh3"], F_INT)
            w, basis, note = _timing(data)
            c = ws[f"{col['Weeks to hearing']}{r}"]
            c.value = w if w is not None else "—"
            c.number_format = F_INT
            c.comment = Comment(note + " Weeks are rounded.", "Parcis", width=260, height=90)
            book.formula(ws, f"{col['Awaiting hearing']}{r}",
                         f"COUNTIFS({lot_rng},{crit},{out_rng},\"Awaiting hearing\")", cnt["pending"], F_INT)
            book.formula(ws, f"{col['Severance attempts']}{r}",
                         f"COUNTIFS({lot_rng},{crit},{type_rng},\"Severance or consent\")",
                         cnt["sever"], F_INT)
        sig = (data.get("stats") or {}).get("signal") or "Sparse"
        c = ws[f"{col['Precedent']}{r}"]
        c.value = SIGNAL_TEXT.get(sig, sig)
        bg, fg = SIGNAL.get(sig, SIGNAL["Sparse"])
        c.fill = _fill(bg)
        c.font = _font(10, True, fg)
        if _unchecked(data, "dev_apps"):
            ws[f"{col['Rezonings / OPAs (OZ)']}{r}"] = "n/a ⚠"
        elif _state(data, "dev_apps") == "skipped":
            ws[f"{col['Rezonings / OPAs (OZ)']}{r}"] = "skipped"
        else:
            tcol = get_column_letter([h for h, _, _ in REZ_COLS].index("Type") + 1)
            dv = f"'{DEV_SHEET}'"
            book.formula(ws, f"{col['Rezonings / OPAs (OZ)']}{r}",
                         f"COUNTIFS({dv}!$A:$A,{crit},{dv}!${tcol}:${tcol},{_q(OZ_NAME)})",
                         sum(1 for d in data.get("devapps") or [] if d.get("type") == "OZ"), F_INT)
        text, warn = _data_check(data)
        c = ws[f"{col['Data check']}{r}"]
        c.value = ("⚠ " if warn else "✓ ") + text
        c.font = _font(10, warn, WARN_INK if warn else "166534")
        if warn:
            c.fill = _fill(WARN_WASH)
        html = (data.get("files") or {}).get("html") or f"{data.get('slug', '')}.html"
        _link(ws[f"{col['Report']}{r}"], html, "Open report")
        if (data.get("links") or {}).get("map"):
            _link(ws[f"{col['Map']}{r}"], data["links"]["map"], "Map")
        if (data.get("links") or {}).get("street_view"):
            _link(ws[f"{col['Street View']}{r}"], data["links"]["street_view"], "Street View")
        for i in range(1, len(SUMMARY_COLS) + 1):
            cell = ws.cell(row=row, column=i)
            cell.border = BOTTOM
            if i >= 4 and SUMMARY_COLS[i - 1][1] not in ("Last decision", "Permit", "Precedent",
                                                         "Data check", "Report", "Map", "Street View"):
                cell.alignment = center
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = _lh(32)

    notes_row = 6 + len(reports) + 1
    notes = [
        "Click an address for its one-page brief, Open report for the interactive map and "
        "tables, or go to the Applications, Permits and Development Applications sheets for every record. "
        "Each record opens the City's own page for it (or its search page, with the number to "
        "paste); the last column of each record sheet opens the raw Open Data row.",
        "Precedent: approved ÷ decided (approved or refused) Committee of Adjustment files "
        "within the radius. Strong ≥ 80%, Mixed 50–79%, Weak < 50%; fewer than 3 decided "
        "files is too few to call.",
        "Counts are formulas over the Applications sheet, so they update if you edit it. "
        "Weeks to hearing is a median; hover the cell for its basis.",
    ]
    for i, t in enumerate(notes):
        _merge(ws, f"A{notes_row + i}:{last}{notes_row + i}", t, _font(9, color=MUTED),
               align=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[notes_row + i].height = _lh(15)
    ws.freeze_panes = "B6"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 110
    ws.sheet_properties.tabColor = SLATE_BLUE
    _print_setup(ws, title_rows="4:5")


# ── one brief per lot ─────────────────────────────────────────────────────────

BRIEF_W = [14, 12, 12, 12, 12, 12, 12, 12, 12, 14]  # A..J
BRIEF_LAST = "J"


def _brief_width(a: str, b: str) -> float:
    return sum(BRIEF_W[_col_index(a) - 1:_col_index(b)])


def _kv(ws, row: int, label: str, value, note: str = "", link: Optional[dict] = None,
        warn: bool = False, value_font: Optional[Font] = None) -> int:
    _merge(ws, f"A{row}:C{row}", label, _font(10, color=INK_2),
           align=Alignment(vertical="top", wrap_text=True))
    text = str(value) if value not in (None, "") else "—"
    if note:
        text += f"\n{note}"
    _merge(ws, f"D{row}:H{row}", text, value_font or _font(10, True, WARN_INK if warn else INK),
           _fill(WARN_WASH) if warn else None, Alignment(vertical="top", wrap_text=True))
    if link and link.get("url"):
        _merge(ws, f"I{row}:J{row}")
        _link(ws[f"I{row}"], link["url"], link.get("text") or "Open")
        ws[f"I{row}"].alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, 11):
        ws.cell(row=row, column=col).border = BOTTOM
    ws.row_dimensions[row].height = max(_lh(18), _lh(13.5) * _lines(text, _brief_width("D", "H")) + 5)
    return row + 1


def _history_events(data: dict) -> list[dict]:
    """The lot's history as dated events (same rules as the HTML timeline)."""
    subj = data.get("subject") or {}
    ev: list[dict] = []
    for c in subj.get("coa") or []:
        ev.append({"date": c.get("filed"), "event": f"{c.get('app') or 'Application'} filed",
                   "file": c.get("ref"), "detail": c.get("desc") or "", "docs": c.get("docs")})
        if c.get("outcome") in ("Approved", "Refused"):
            kind = "Consent" if c.get("app") == "Consent" else "Variance"
            bits = []
            if c.get("days") is not None:
                bits.append(f"{c['days']} days after filing.")
            if c.get("appeal"):
                bits.append(f"Appeal: {c['appeal']}.")
            ev.append({"date": c.get("hearing"),
                       "event": f"{kind} {c['outcome'].lower()} by the Committee of Adjustment",
                       "file": c.get("ref"), "detail": " ".join(bits), "docs": c.get("docs"),
                       "bad": c["outcome"] == "Refused"})
        elif c.get("hearing") or c.get("outcome") == "Pending":
            ev.append({"date": c.get("hearing") or c.get("filed"), "event": c.get("label"),
                       "file": c.get("ref"), "detail": "", "docs": c.get("docs")})
    for p in subj.get("permits") or []:
        noun = {"New building": "New-building permit", "Demolition": "Demolition permit"}.get(
            p.get("category"), "Building permit")
        desc = re.sub(r"^[A-Za-z()]+\s+-\s+", "", p.get("desc") or "").split(" See also")[0]
        extra = []
        if not p.get("cost_unreliable") and p.get("cost"):
            extra.append(f"Estimated cost ${p['cost']:,}.")
        if p.get("linked_coa"):
            extra.append(f"Cites variance {', '.join(p['linked_coa'])}.")
        if p.get("applied"):
            ev.append({"date": p["applied"], "event": f"{noun} applied for", "file": p.get("root"),
                       "detail": desc, "docs": p.get("docs")})
        if p.get("issued"):
            ev.append({"date": p["issued"], "event": f"{noun} issued", "file": p.get("root"),
                       "detail": " ".join(extra), "docs": p.get("docs")})
        if p.get("completed"):
            ev.append({"date": p["completed"], "event": f"{noun} closed", "file": p.get("root"),
                       "detail": "", "docs": p.get("docs")})
    ev = [e for e in ev if e.get("date")]
    ev.sort(key=lambda e: e["date"])
    ps = subj.get("permits") or []
    open_p = next((p for p in ps if p.get("category") == "New building" and not p.get("completed")),
                  next((p for p in ps if not p.get("completed")), None))
    if open_p:
        ev.append({"date": data.get("generated"), "event": "Status at this run",
                   "file": open_p.get("root"),
                   "detail": f"Permit {open_p.get('root')}: {open_p.get('status_text', '')}.",
                   "docs": open_p.get("docs"), "now": True})
    return ev


def _write_brief(book: _Book, ws, data: dict, label: str) -> None:
    for i, w in enumerate(BRIEF_W, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = ZONE
    zone = data.get("zone") or {}
    env = data.get("envelope") or {}
    stats = data.get("stats") or {}
    links = data.get("links") or {}
    radius = data.get("radius_m") or DEFAULT_RADIUS_M

    # 1-3 title, meta, navigation
    _merge(ws, "A1:H1", label, _font(18, True, WHITE), _fill(SLATE_BLUE), Alignment(vertical="center", indent=1))
    _merge(ws, "I1:J1", zone.get("code") or "—", _font(14, True, ZONE_INK), _fill(ZONE),
           Alignment(horizontal="center", vertical="center"))
    ws.row_dimensions[1].height = 38
    meta = "  ·  ".join(x for x in [
        (data.get("ward") or "").replace("-", "–"), data.get("fsa") and f"Postal area {data['fsa']}",
        data.get("district") and f"District {data['district']}", f"{radius} m radius",
        f"Run {_fmt_date(data.get('generated'))}"] if x)
    meta = f"{meta}  ·  {NOTICE}"
    _merge(ws, "A2:J2", meta, _font(10, color=INK_2), _fill(CREAM),
           Alignment(vertical="center", indent=1, wrap_text=True))
    ws.row_dimensions[2].height = max(_lh(22), _lh(13) * _lines(meta, sum(BRIEF_W)) + 8)
    _link(ws["A3"], _ref("Summary"), "← Summary", internal=True)
    html = (data.get("files") or {}).get("html") or f"{data.get('slug', '')}.html"
    _merge(ws, "B3:C3")
    _link(ws["B3"], html, "Interactive report")
    if links.get("map"):
        _link(ws["D3"], links["map"], "Google Maps")
    if links.get("street_view"):
        _link(ws["E3"], links["street_view"], "Street View")
    if links.get("zoning_map"):
        _merge(ws, "F3:G3")
        _link(ws["F3"], links["zoning_map"], "City zoning map")
    _merge(ws, "H3:J3")
    _link(ws["H3"], _ref("Applications", "A4"), f"All {len(_records(data))} applications →", internal=True)
    ws.row_dimensions[3].height = _lh(20)

    # bottom line
    row = _section(ws, 5, BRIEF_LAST, "Bottom line")
    sig = stats.get("signal") or "Sparse"
    cnt = _counts(data)
    if sig == "Unchecked":
        verdict = "Neighbour precedent couldn't be checked: the City's data wasn't reachable for this run."
    elif cnt["decided"]:
        verdict = (f"{SIGNAL_TEXT.get(sig, sig)}: {cnt['approved']} of {cnt['decided']} decided "
                   f"applications within {radius} m were approved ({zr.round_half_up(cnt['rate'] * 100)}%).")
    elif cnt["apps"]:
        verdict = f"{_plural(cnt['apps'], 'application is', 'applications are')} on file within {radius} m, none decided yet."
    else:
        verdict = f"No Committee of Adjustment applications on record within {radius} m."
    bg, fg = SIGNAL.get(sig, SIGNAL["Sparse"])
    _merge(ws, f"A{row}:J{row}", verdict, _font(12, True, fg), _fill(bg),
           Alignment(vertical="center", indent=1, wrap_text=True))
    ws.row_dimensions[row].height = _lh(26)
    row += 1
    summary = data.get("summary") or {}
    text = summary.get("text") or ""
    _merge(ws, f"A{row}:J{row}", text, _font(11, color=INK), align=Alignment(wrap_text=True, vertical="top"))
    ws.row_dimensions[row].height = _lh(14.5) * (_lines(text, sum(BRIEF_W) * 0.93) + 0.4) + 4
    row += 1
    src = ("Written by Claude (" + (summary.get("model_version") or "") + ") from the computed facts; "
           "every number in it was checked against them before use."
           if summary.get("source") == "claude"
           else "Written from the computed facts by a fixed template.")
    _merge(ws, f"A{row}:J{row}", src, _font(9, color=MUTED, italic=True), align=Alignment(vertical="top"))
    row += 2

    # key numbers (label / value / basis)
    row = _section(ws, row, BRIEF_LAST, "Key numbers")
    lot_rng = f"Applications!${_COL['Lot']}:${_COL['Lot']}"
    out_rng = f"Applications!${_COL['Outcome']}:${_COL['Outcome']}"
    type_rng = f"Applications!${_COL['Project type']}:${_COL['Project type']}"
    st_rng = f"Applications!${_COL['Storeys']}:${_COL['Storeys']}"
    crit = _q(label)
    w, basis, note = _timing(data)
    unchecked = _unchecked(data, "neighbour_coa")
    tiles = [
        ("A", "B", f"Applications within {radius} m", f"COUNTIFS({lot_rng},{crit})", cnt["apps"],
         F_INT, "Committee of Adjustment files"),
        ("C", "D", "Approval rate",
         f"IFERROR(COUNTIFS({lot_rng},{crit},{out_rng},\"Approved\")/(COUNTIFS({lot_rng},{crit},{out_rng},\"Approved\")"
         f"+COUNTIFS({lot_rng},{crit},{out_rng},\"Refused\")),\"–\")", cnt["rate"] if cnt["decided"] else "–",
         F_PCT, f"{cnt['approved']} of {cnt['decided']} decided"),
        ("E", "F", "Weeks to hearing", None, w if w is not None else "–", F_INT, basis),
        ("G", "H", "New houses approved",
         f"COUNTIFS({lot_rng},{crit},{type_rng},\"New house\",{out_rng},\"Approved\")",
         _storey_counts(data)["all"], F_INT, f"{cnt['nh3']} of them three storeys"),
        ("I", "J", "Awaiting hearing", f"COUNTIFS({lot_rng},{crit},{out_rng},\"Awaiting hearing\")",
         cnt["pending"], F_INT, "filed, not yet heard"),
    ]
    for a, b, lab, f, val, fmt, sub in tiles:
        _merge(ws, f"{a}{row}:{b}{row}", lab, _font(9, color=INK_2), align=Alignment(vertical="bottom", wrap_text=True))
        ws.merge_cells(f"{a}{row + 1}:{b}{row + 1}")
        cell = f"{a}{row + 1}"
        big = _font(20, True, INK)
        if unchecked:
            ws[cell] = "n/a"
            ws[cell].font = _font(20, True, WARN_INK)
        elif f:
            book.formula(ws, cell, f, val, fmt, big, Alignment(horizontal="left", vertical="center"))
        else:
            ws[cell] = val
            ws[cell].number_format = fmt
            ws[cell].font = big
            ws[cell].alignment = Alignment(horizontal="left", vertical="center")
            ws[cell].comment = Comment(note + " Weeks are rounded.", "Parcis", width=260, height=90)
        _merge(ws, f"{a}{row + 2}:{b}{row + 2}", "couldn't be checked" if unchecked else sub,
               _font(9, color=MUTED), align=Alignment(vertical="top", wrap_text=True))
    ws.row_dimensions[row].height = _lh(16)
    ws.row_dimensions[row + 1].height = 30
    ws.row_dimensions[row + 2].height = _lh(28)
    row += 4

    # as of right
    row = _section(ws, row, BRIEF_LAST, "What you can build as-of-right",
                   "Zone string decoded, plus the overlays and plans on top of it. The decoder reads "
                   "the label only; read the by-law and any exception before design.")
    zl = zone.get("links") or {}
    if zone.get("raw"):
        row = _kv(ws, row, "Zone", zone["raw"], zone.get("name", "") and f"{zone['name']} zone",
                  {"url": (zl.get("chapter") or {}).get("url"), "text": "By-law chapter"},
                  value_font=_font(11, True, INK))
        for t in zone.get("tokens") or []:
            if t.get("key") == "zone":
                continue
            lab = f"   {t.get('tok')}"
            if t.get("key") == "exception":
                row = _kv(ws, row, lab, t.get("label", ""),
                          "Can change the standard rules; read it before design.",
                          {"url": (zl.get("exception") or {}).get("url"), "text": "Exception text"})
            elif t.get("key") == "fsi":
                row = _kv(ws, row, lab, t.get("label", ""),
                          f"Gross floor area up to {t.get('value')} × lot area.")
            elif t.get("key") == "frontage_m":
                row = _kv(ws, row, lab, t.get("label", ""), "Matters for any severance.")
            else:
                row = _kv(ws, row, lab, t.get("label") or f"{t.get('tok')}: not decoded")
    else:
        row = _kv(ws, row, "Zone", "Couldn't be checked" if _unchecked(data, "zone_map")
                  else "Not on the City-wide zoning map",
                  "" if _unchecked(data, "zone_map") else
                  "A former municipal by-law probably applies; check with the City.",
                  warn=_unchecked(data, "zone_map"))

    def layer(key, found, none_text, found_note=""):
        st = _state(data, key)
        if st == "unchecked":
            return "Couldn't be checked", "The layer failed to load for this run.", True
        if st == "skipped":
            return "Not checked", "Skipped for this run.", False
        return (found, found_note, False) if found else (none_text, "", False)

    h = env.get("height_m")
    v, n_, wn = layer("height", h and f"{h} m", "No height overlay at this lot",
                      env.get("height_string") and f"Height overlay {env['height_string']}")
    row = _kv(ws, row, "Maximum height", v, n_, warn=wn)
    cov = env.get("coverage_pct")
    v, n_, wn = layer("coverage", cov and f"{cov}%", "No lot coverage overlay at this lot")
    if not cov and _state(data, "coverage") == "none" and env.get("coverage_nearest_m") is not None:
        n_ = f"Nearest coverage area is about {env['coverage_nearest_m'] / 1000:.1f} km away."
    row = _kv(ws, row, "Lot coverage overlay", v, n_, warn=wn)
    v, n_, wn = layer("setback", env.get("setback_label"), "None at this lot")
    row = _kv(ws, row, "Building setback overlay", v, n_, warn=wn)
    v, n_, wn = layer("secondary_plan", env.get("secondary_plan"), "None")
    row = _kv(ws, row, "Secondary plan", v, n_, warn=wn)
    sasp = env.get("area_specific_policy")
    v, n_, wn = layer("area_specific", sasp and f"SASP {sasp}", "None")
    row = _kv(ws, row, "Site and area specific policy", v, n_, warn=wn)
    zsrc = {"zoning_map": "City zoning map", "coa_designation": "This lot's variance file",
            "outside_zoning_map": "Not on the By-law 569-2013 map", "unresolved": "Not resolved"}.get(
        zone.get("source"), zone.get("source") or "—")
    last_ref = next((c.get("ref") for c in (data.get("subject") or {}).get("coa") or []), "")
    zn = ""
    if zone.get("source") == "zoning_map" and zone.get("coa_string"):
        zn = f"Matches variance file {last_ref}".strip() if zone.get("matches_coa") \
            else f"Variance file lists {zone['coa_string']}"
    row = _kv(ws, row, "Zone source", zsrc, zn)
    row += 1

    # the lot's history
    row = _section(ws, row, BRIEF_LAST, "This lot's history",
                   (data.get("subject") or {}).get("chain") or "")
    heads = [("A", "A", "Date"), ("B", "D", "What happened"), ("E", "F", "File or permit"),
             ("G", "I", "Detail"), ("J", "J", "Source")]
    for a, b, h_ in heads:
        _merge(ws, f"{a}{row}:{b}{row}", h_, _font(9, True, WHITE), _fill(SLATE), Alignment(vertical="center", indent=0 if a == "J" else 1))
    ws.row_dimensions[row].height = _lh(18)
    row += 1
    events = _history_events(data)
    if _unchecked(data, "subject_coa") and _unchecked(data, "subject_permits"):
        _merge(ws, f"A{row}:J{row}", "This lot's history couldn't be checked: the City's data wasn't "
               "reachable for this run. This is missing data, not a result.",
               _font(10, True, WARN_INK), _fill(WARN_WASH), Alignment(wrap_text=True, indent=1))
        row += 2
    elif not events:
        _merge(ws, f"A{row}:J{row}", "No Committee of Adjustment file or building permit is recorded "
               "for this address since 2017.", _font(10, color=INK_2), align=Alignment(indent=1))
        row += 2
    else:
        for e in events:
            d = _date(e["date"])
            ws[f"A{row}"] = d
            ws[f"A{row}"].number_format = F_DATE
            ws[f"A{row}"].font = _font(10, color=INK_2)
            ws[f"A{row}"].alignment = Alignment(horizontal="left", vertical="top", indent=1)
            _merge(ws, f"B{row}:D{row}", e["event"], _font(10, True, "991B1B" if e.get("bad") else
                                                             (LINK if e.get("now") else INK)),
                   align=Alignment(vertical="top", wrap_text=True, indent=1))
            _merge(ws, f"E{row}:F{row}", e.get("file") or "", _font(10, color=INK_2),
                   align=Alignment(vertical="top", indent=1))
            _merge(ws, f"G{row}:I{row}", e.get("detail") or "", _font(9, color=INK_2),
                   align=Alignment(vertical="top", wrap_text=True, indent=1))
            if doc_text(e.get("docs")):
                _link(ws[f"J{row}"], doc_route(e["docs"])["url"], doc_text(e["docs"]))
                ws[f"J{row}"].alignment = Alignment(vertical="top", wrap_text=True)
            for col in range(1, 11):
                ws.cell(row=row, column=col).border = BOTTOM
            ws.row_dimensions[row].height = max(_lh(18), _lh(12.5) * max(
                _lines(e["event"], _brief_width("B", "D")), _lines(e.get("detail"), _brief_width("G", "I") * 1.1)) + 5)
            row += 1
        row += 1

    # precedent by project type
    row = _section(ws, row, BRIEF_LAST, f"What the neighbours got approved within {radius} m",
                   "Counts are live formulas over the Applications sheet. Approval rate counts "
                   "approved and refused decisions only.")
    heads = ["Project type", "", "", "Applications", "Approved", "Refused", "Awaiting hearing",
             "Not decided", "Approval rate", ""]
    _merge(ws, f"A{row}:C{row}", "Project type", _font(9, True, WHITE), _fill(SLATE), Alignment(vertical="center", indent=1))
    for i, h_ in enumerate(heads[3:9], start=4):
        c = ws.cell(row=row, column=i, value=h_)
        c.font = _font(9, True, WHITE)
        c.fill = _fill(SLATE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=row, column=10).fill = _fill(SLATE)
    ws.row_dimensions[row].height = _lh(28)
    row += 1
    if unchecked:
        _merge(ws, f"A{row}:J{row}", "Neighbour applications couldn't be checked for this run. "
               "This is missing data, not a result.", _font(10, True, WARN_INK), _fill(WARN_WASH),
               Alignment(indent=1))
        row += 2
    else:
        first = row
        for t in PROJECT_TYPES:
            tc = _type_counts(data, t)
            if not tc["n"]:
                continue
            _merge(ws, f"A{row}:C{row}", t, _font(10), align=Alignment(vertical="center", indent=1))
            tcrit = _q(t)
            base = f"{lot_rng},{crit},{type_rng},{tcrit}"
            book.formula(ws, f"D{row}", f"COUNTIFS({base})", tc["n"], F_INT, _font(10, True),
                         Alignment(horizontal="center"))
            for col, key, lab in (("E", "Approved", "Approved"), ("F", "Refused", "Refused"),
                                  ("G", "Pending", "Awaiting hearing"), ("H", "Other", "Not decided")):
                book.formula(ws, f"{col}{row}", f"COUNTIFS({base},{out_rng},{_q(lab)})", tc[key], F_INT,
                             align=Alignment(horizontal="center"))
            book.formula(ws, f"I{row}", f"IF(E{row}+F{row}>0,E{row}/(E{row}+F{row}),\"–\")",
                         tc["rate"] if tc["rate"] != "" else "–", F_PCT, _font(10, True),
                         Alignment(horizontal="center"))
            for col in range(1, 11):
                ws.cell(row=row, column=col).border = BOTTOM
            row += 1
        if row == first:
            _merge(ws, f"A{row}:J{row}", f"No Committee of Adjustment applications on record within "
                   f"{radius} m.", _font(10, color=INK_2), align=Alignment(indent=1))
            row += 1
        row += 1
        # storeys of approved new houses
        sc = _storey_counts(data)
        _merge(ws, f"A{row}:C{row}", "Approved new houses by storeys", _font(10, True, INK_2),
               align=Alignment(vertical="center", indent=1))
        for i, (s, lab) in enumerate(((3, "3 storeys"), (2.5, "2½ storeys"), (2, "2 storeys"))):
            colL = get_column_letter(4 + i * 2)
            colV = get_column_letter(5 + i * 2)
            ws[f"{colL}{row}"] = lab
            ws[f"{colL}{row}"].font = _font(9, color=INK_2)
            ws[f"{colL}{row}"].alignment = Alignment(horizontal="right", vertical="center")
            book.formula(ws, f"{colV}{row}",
                         f"COUNTIFS({lot_rng},{crit},{type_rng},\"New house\",{out_rng},\"Approved\",{st_rng},{s})",
                         sc[s], F_INT, _font(12, True), Alignment(horizontal="left", vertical="center", indent=1))
        rest = sc["other"] + sc["unknown"]
        if rest:
            _merge(ws, f"J{row}:J{row}", f"+{rest} other", _font(9, color=MUTED), align=Alignment(vertical="center"))
        ws.row_dimensions[row].height = _lh(22)
        row += 2

    # signs of change
    row = _section(ws, row, BRIEF_LAST, "Signs of change",
                   "Neighbours' activity that could shift what gets approved here next.")
    recs = _records(data)
    items = [("Awaiting hearing", r, f"{r.get('type')}{', ' + _storey_word(r.get('storeys')) if r.get('storeys') else ''}. "
              f"Filed {_fmt_date(r.get('filed'))}, {r.get('ref')}.", r.get("docs"))
             for r in recs if r.get("outcome") == "Pending"]
    items += [("Severance attempt", r, f"{(r.get('desc') or 'No description').split('. ')[0].rstrip('.')}. "
               f"Filed {_fmt_date(r.get('filed'))}; {(r.get('label') or '').lower()}.", r.get("docs"))
              for r in recs if r.get("type") == "Severance or consent"]
    items += [("Development application", d, f"{d.get('type_name') or d.get('type')}, {d.get('status')}. {d.get('no')}"
               f"{', submitted ' + _fmt_date(d.get('submitted')) if d.get('submitted') else ''}. "
               f"{d.get('desc') or ''}", d.get("docs"))
              for d in data.get("devapps") or [] if devapp_is_sign(d, data.get("generated"))]
    if _unchecked(data, "neighbour_coa") or _unchecked(data, "dev_apps"):
        _merge(ws, f"A{row}:J{row}", "Some of this couldn't be checked for this run (City data "
               "unreachable).", _font(10, True, WARN_INK), _fill(WARN_WASH), Alignment(indent=1))
        row += 1
    if not items:
        extra = ""
        if data.get("fsa_devapps") is not None:
            extra = (f" The wider {data.get('fsa')} postal area has "
                     f"{data['fsa_devapps']} development applications on file.")
        _merge(ws, f"A{row}:J{row}", f"Nothing pending, no severance attempts and no open or recent development applications within "
               f"{radius} m.{extra}", _font(10, color=INK_2), align=Alignment(wrap_text=True, indent=1))
        ws.row_dimensions[row].height = _lh(28)
        row += 1
    for kind, r, text, docs in items:
        ws[f"A{row}"] = kind
        ws[f"A{row}"].font = _font(9, True, INK_2)
        ws[f"A{row}"].alignment = Alignment(vertical="top", wrap_text=True, indent=1)
        _merge(ws, f"B{row}:D{row}", f"{r.get('addr')} · {r.get('d')} m", _font(10, True),
               align=Alignment(vertical="top", wrap_text=True))
        _merge(ws, f"E{row}:I{row}", text, _font(9, color=INK_2), align=Alignment(vertical="top", wrap_text=True))
        if doc_text(docs):
            _link(ws[f"J{row}"], doc_route(docs)["url"], doc_text(docs))
            ws[f"J{row}"].alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, 11):
            ws.cell(row=row, column=col).border = BOTTOM
        ws.row_dimensions[row].height = max(
            _lh(18), _lh(12.5) * max(_lines(text, _brief_width("E", "I") * 1.1),
                                     _lines(kind, BRIEF_W[0] * 0.8),
                                     _lines(f"{r.get('addr')} · {r.get('d')} m", _brief_width("B", "D")),
                                     _lines(doc_text(docs), BRIEF_W[9] * 0.9)) + 5)
        row += 1
    row += 1
    _merge(ws, f"A{row}:J{row}", "Contains information licensed under the Open Government Licence – "
           "Toronto. Not a legal opinion; confirm with the City before relying on it.",
           _font(8, color=MUTED), align=Alignment(wrap_text=True))
    # Every text label in column A wraps (merged or not), so none is cut off.
    for (cell,) in ws.iter_rows(min_row=4, max_row=row, max_col=1):
        if isinstance(cell.value, str):
            al = cell.alignment
            cell.alignment = Alignment(horizontal=al.horizontal, vertical=al.vertical or "top",
                                       indent=al.indent, wrap_text=True)
    _print_setup(ws)
    ws.print_area = f"A1:J{row}"


def _storey_word(s) -> str:
    return {3: "3 storeys", 2.5: "2½ storeys", 2: "2 storeys", 1: "1 storey", 4: "4 storeys"}.get(s, f"{s} storeys")


# ── record sheets ─────────────────────────────────────────────────────────────

def _record_sheet(ws, title: str, subtitle: str, cols: list, rows: list, table: str,
                  style: dict) -> None:
    last = get_column_letter(len(cols))
    _title_band(ws, last, title, subtitle)
    _link(ws["A3"], _ref("Summary"), "← Summary", internal=True)
    ws.row_dimensions[3].height = 20
    for i, (h, w, _) in enumerate(cols, start=1):
        c = ws.cell(row=_TABLE_TOP, column=i, value=h)
        c.font = _font(9, True, WHITE)
        c.fill = _fill(SLATE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[_TABLE_TOP].height = 30
    for n, values in enumerate(rows, start=_TABLE_TOP + 1):
        wrap_lines = 1
        for i, (h, w, fmt) in enumerate(cols, start=1):
            v = values.get(h)
            c = ws.cell(row=n, column=i)
            if isinstance(v, tuple) and v and v[0] == "link":
                if v[1]:
                    _link(c, v[1], v[2])
                c.alignment = Alignment(vertical="top")
                continue
            c.value = v
            c.font = _font(10)
            if fmt:
                c.number_format = fmt
            c.alignment = (Alignment(horizontal="center", vertical="center") if h == "Dist. (m)"
                           else Alignment(vertical="top", wrap_text=w >= 40))
            if w >= 40 and v:
                wrap_lines = max(wrap_lines, _lines(str(v), w))
            if h in style and v in style[h]:
                bg, fg = style[h][v]
                c.fill = _fill(bg)
                c.font = _font(10, True, fg)
        ws.row_dimensions[n].height = min(96, max(15, 13 * min(wrap_lines, 7) + 3))
        for c in ws[n]:
            c.border = BOTTOM
    if rows:
        ref = f"A{_TABLE_TOP}:{last}{_TABLE_TOP + len(rows)}"
        t = Table(displayName=table, ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
        ws.add_table(t)
    else:
        ws.cell(row=_TABLE_TOP + 1, column=1, value="No records within any lot's radius.").font = _font(10, color=INK_2)
    ws.freeze_panes = f"D{_TABLE_TOP + 1}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = STONE
    _print_setup(ws, title_rows=f"{_TABLE_TOP}:{_TABLE_TOP}")


def _application_rows(reports: list) -> list:
    rows = []
    for data in reports:
        if _unchecked(data, "neighbour_coa"):
            continue
        lab = lot_label(data)
        for r in _records(data):
            rows.append({
                "Lot": lab, "Dist. (m)": r.get("d"), "Address": r.get("addr"),
                "File": r.get("ref") or "No file number", "Documents": _doc(r.get("docs")),
                "Application": r.get("app"), "Project type": r.get("type"),
                "Storeys": r.get("storeys"), "Outcome": OUTCOME.get(r.get("outcome"), ("Not decided",))[0],
                "Decision on file": r.get("label"), "Filed": _date(r.get("filed")),
                "Hearing": _date(r.get("hearing")),
                "Weeks to hearing": zr.weeks(r["days"]) if r.get("days") is not None else None,
                "Appeal": r.get("appeal") or None,
                "Built under permit": ", ".join(r.get("permits") or []) or None,
                "What was asked": r.get("desc") or "No description on file",
                "Map": ("link", (r.get("links") or {}).get("map"), "Map"),
                "Raw City data": ("link", r.get("src"), "Raw data"),
            })
    return rows


def _permit_rows(reports: list) -> list:
    rows = []
    for data in reports:
        if _unchecked(data, "neighbour_permits"):
            continue
        lab = lot_label(data)
        for p in data.get("permits") or []:
            note = None
            if p.get("cost") is None:
                note = "Not declared"
            elif p.get("cost_unreliable"):
                note = "Unreliable (zero or out of band)"
            rows.append({
                "Lot": lab, "Dist. (m)": p.get("d"), "Address": p.get("addr"), "Permit": p.get("root"),
                "Documents": _doc(p.get("docs")),
                "Category": p.get("category"), "Permit types": "; ".join(p.get("types") or []),
                "Structure": "; ".join(p.get("structure") or []),
                "Status": p.get("status_text") or p.get("status"),
                "Applied": _date(p.get("applied")), "Issued": _date(p.get("issued")),
                "Completed": _date(p.get("completed")),
                "Est. cost": p.get("cost") if p.get("cost") and not p.get("cost_unreliable") else None,
                "Cost note": note, "Units": _num(p.get("units")) or None,
                "Cites variance": ", ".join(p.get("linked_coa") or []) or None,
                "Description": re.sub(r"\s+", " ", p.get("desc") or "").strip() or None,
                "Map": ("link", (p.get("links") or {}).get("map"), "Map"),
                "Raw City data": ("link", p.get("src"), "Raw data"),
            })
    return rows


def _rezoning_rows(reports: list) -> list:
    rows = []
    for data in reports:
        if _unchecked(data, "dev_apps"):
            continue
        lab = lot_label(data)
        for d in data.get("devapps") or []:
            rows.append({
                "Lot": lab, "Dist. (m)": d.get("d"), "Address": d.get("addr"), "Application": d.get("no"),
                "Documents": _doc(d.get("docs")),
                "Type": d.get("type_name") or d.get("type"), "Status": d.get("status"), "Submitted": _date(d.get("submitted")),
                "Description": d.get("desc") or None,
                "Map": ("link", (d.get("links") or {}).get("map"), "Map"),
                "Raw City data": ("link", d.get("src"), "Raw data"),
            })
    return rows


# ── Sources & method ──────────────────────────────────────────────────────────

def _source_row(ws, row: int, title: str, used_for: str, url: str, link_text: str,
                widths: list) -> int:
    a = ws.cell(row=row, column=1, value=title)
    a.font = _font(10, True)
    a.alignment = Alignment(wrap_text=True, vertical="top")
    b = ws.cell(row=row, column=2, value=used_for)
    b.font = _font(10, color=INK_2)
    b.alignment = Alignment(wrap_text=True, vertical="top")
    if url:
        _link(ws.cell(row=row, column=3), url, link_text)
        ws.cell(row=row, column=3).alignment = Alignment(vertical="top")
    for col in range(1, 4):
        ws.cell(row=row, column=col).border = BOTTOM
    ws.row_dimensions[row].height = 13 * max(_lines(title, widths[0]), _lines(used_for, widths[1])) + 5
    return row + 1


def _write_sources(ws, reports: list) -> None:
    widths = [36, 62, 16] + [21] * max(1, len(reports))
    last = get_column_letter(len(widths))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _title_band(ws, last, "Sources and method",
                "Where every number comes from, how it's counted, and what this workbook can't tell you.")
    _link(ws["A3"], _ref("Summary"), "← Summary", internal=True)
    row = _section(ws, 5, last, "City of Toronto Open Data used")
    for i, h in enumerate(("Dataset", "What it's used for", "Link"), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = _font(9, True, WHITE)
        c.fill = _fill(SLATE)
        c.alignment = Alignment(vertical="center", indent=0)
    ws.row_dimensions[row].height = 18
    row += 1
    ds = next((d.get("datasets") for d in reports if d.get("datasets")), None) or []
    for d in ds:
        row = _source_row(ws, row, d.get("title"), d.get("used_for"), d.get("url"), "open.toronto.ca",
                          widths)
    L0 = reports[0].get("links", {}) if reports else {}
    extra = [("Application Information Centre", "Drawings, decisions and documents for open and "
              "recently final Committee of Adjustment files and development applications.", L0.get("aic")),
             ("Committee of Adjustment staff", "Decision notices for files the AIC no longer lists, "
              "and anything older than 10 years.", L0.get("coa_staff")),
             ("Research Request Portal", "10 years of CoA decision notices within 500 m ($150 + HST) "
              "or 1,000 m ($300 + HST) of an address.", L0.get("research_request")),
             ("Building Application Status", "Status and inspections for open permits up to 10 years "
              "from application, and permits closed in the last month.", L0.get("permit_status")),
             ("Building records request", "Drawings and records for older permits: "
              "bldrecords@toronto.ca, $76.98, up to 30 business days.", L0.get("building_records")),
             ("Zoning By-law 569-2013", "The by-law text, by chapter and exception.",
              reports[0].get("links", {}).get("bylaw") if reports else ""),
             ("Interactive zoning map", "The City's own zoning map, centred on the lot.",
              L0.get("zoning_map"))]
    for t, u, url in extra:
        row = _source_row(ws, row, t, u, url, "toronto.ca", widths)
    row += 1

    row = _section(ws, row, last, "Data check for each run",
                   "\"Couldn't be checked\" means the City's service wasn't reachable. It is missing "
                   "data, never a finding that there are no records.")
    ws.cell(row=row, column=1, value="Source").font = _font(9, True, WHITE)
    ws.cell(row=row, column=1).fill = _fill(SLATE)
    for col in (2, 3):
        ws.cell(row=row, column=col).fill = _fill(SLATE)
    for j, d in enumerate(reports, start=4):
        c = ws.cell(row=row, column=j, value=f"{lot_label(d)}\n{d.get('generated_at') or d.get('generated', '')}")
        c.font = _font(9, True, WHITE)
        c.fill = _fill(SLATE)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 30
    row += 1
    names = next((d.get("source_names") for d in reports if d.get("source_names")), {}) or {}
    for key in names:
        _merge(ws, f"A{row}:C{row}", names[key][:1].upper() + names[key][1:], _font(10))
        for j, d in enumerate(reports, start=4):
            st = _state(d, key)
            c = ws.cell(row=row, column=j, value=STATE_TEXT.get(st, st or "—") + (" ⚠" if st == "unchecked" else ""))
            c.font = _font(10, st == "unchecked", WARN_INK if st == "unchecked" else (INK if st == "found" else INK_2))
            if st == "unchecked":
                c.fill = _fill(WARN_WASH)
        for col in range(1, len(widths) + 1):
            ws.cell(row=row, column=col).border = BOTTOM
        row += 1
    row += 1

    row = _section(ws, row, last, "How the numbers are made")
    rules = [
        ("Decided", "Approved or refused. Deferred, withdrawn, awaiting-hearing and closed-without-a-"
                    "decision files are listed but not counted."),
        ("Approval rate", "Approved ÷ decided, for Committee of Adjustment files within the radius."),
        ("Precedent", "Strong: 80% or more approved. Mixed: 50–79%. Weak: under 50%. Too few decisions: "
                      "fewer than 3 decided files."),
        ("Weeks to hearing", "Median time from filing to hearing for decided files; \"last 3 yrs\" counts "
                             "files filed in the three years before the run. Shown once 3 or more files qualify."),
        ("Project type", "From the City's application sub-type (new house or addition), refined by keywords "
                         "for pools and decks, legalizations and severances."),
        ("Storeys", "Read from the first storey phrase in a new-house description."),
        ("Built under permit", "A variance counts as built when a building permit's description cites its "
                               "decision number. That makes it a floor, not a full count."),
        ("Permit project", "Permits grouped by permit-number root (for example 22 111410 BLD, PLB, DRN and "
                           "HVA are one project)."),
        ("Est. cost", "The declared construction value on the permit. Left blank when missing, zero or "
                      "outside the usual range for its structure type."),
        ("Distance", "Straight line between the City's address points, in metres."),
        ("Documents", "Where to get the record's papers, the same route as the HTML record card: "
                      "Open on City site (the AIC still carries the file); Request from CoA staff "
                      "(the AIC drops CoA files about 90 days after they're final; the Research "
                      "Request Portal also sells 10 years of decision notices within 500 m for $150 "
                      "+ HST or 1,000 m for $300 + HST); Check permit status (open permits up to 10 "
                      "years from application, closed ones for a month); Building records request "
                      "(bldrecords@toronto.ca, $76.98, up to 30 business days); or "
                      "developmentreview@toronto.ca for closed development applications."),
        ("Raw City data", "Opens the exact Open Data row(s) the numbers came from, as raw JSON."),
    ]
    for k, v in rules:
        ws.cell(row=row, column=1, value=k).font = _font(10, True)
        ws.cell(row=row, column=1).alignment = Alignment(vertical="top")
        _merge(ws, f"B{row}:{last}{row}", v, _font(10, color=INK_2), align=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[row].height = max(16, 13 * _lines(v, sum(widths[1:])) + 3)
        row += 1
    row += 1
    row = _section(ws, row, last, "Limits")
    limits = [
        "Precedent covers applications in each lot's postal area only, so files just across its boundary "
        "are missing.",
        "The City's open data describes each project but not which by-law rules were varied or by how much, "
        "so requested and permitted numbers can't be compared.",
        "Exception and by-law text isn't read, and the Official Plan land-use designation isn't loaded.",
        "Records start in 2017 (closed CoA files and cleared permits).",
    ]
    for t in limits:
        _merge(ws, f"A{row}:{last}{row}", "•  " + t, _font(10, color=INK_2), align=Alignment(wrap_text=True, vertical="top"))
        ws.row_dimensions[row].height = max(16, 13 * _lines(t, sum(widths)) + 3)
        row += 1
    row += 1
    _merge(ws, f"A{row}:B{row}", "Contains information licensed under the Open Government Licence – Toronto.",
           _font(9, color=MUTED))
    _link(ws[f"C{row}"], LICENCE_URL, "Licence")
    row += 1
    _merge(ws, f"A{row}:{last}{row}", "Parcis Zoning Agent. Not a legal opinion; confirm with the City before "
           "relying on it.", _font(9, color=MUTED))
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = SAND
    _print_setup(ws)


# ── cached values for formulas ────────────────────────────────────────────────

def _xml_value(v) -> tuple[str, str]:
    """(type attribute, <v> text) for a cached formula result."""
    if isinstance(v, bool):
        return ' t="b"', "1" if v else "0"
    if isinstance(v, (int, float)):
        return "", repr(float(v)) if isinstance(v, float) else str(v)
    s = str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    return ' t="str"', s


def _inject_cached_values(path: Path, cache: dict, sheet_files: dict) -> None:
    """Write each formula's value into the saved file (openpyxl can't), so viewers that
    don't recalculate show numbers. Excel recalculates on open anyway (fullCalcOnLoad)."""
    if not cache:
        return
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            title = sheet_files.get(item.filename)
            if title in cache:
                xml = data.decode("utf-8")
                for coord, value in cache[title].items():
                    t, v = _xml_value(value)
                    pat = re.compile(r'<c r="' + coord + r'"((?: [a-z]+="[^"]*")*)><f>(.*?)</f><v\s*/>(?:</v>)?</c>')
                    xml = pat.sub(lambda m: f'<c r="{coord}"{m.group(1)}{t}><f>{m.group(2)}</f><v>{v}</v></c>', xml, count=1)
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


# ── build ─────────────────────────────────────────────────────────────────────

def load_reports(report_dir: Optional[Path] = None) -> list[dict]:
    """Every report JSON in the folder, sorted by address then radius."""
    report_dir = Path(report_dir or REPORT_DIR)
    out = []
    for p in sorted(report_dir.glob("*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, ValueError) as e:
            print(f"  [warn] skipped {p.name}: {e}")
            continue
        if isinstance(d, dict) and str(d.get("schema_version", "")).startswith("1.") and d.get("address"):
            d.setdefault("files", {}).setdefault("html", f"{p.stem}.html")
            out.append(d)
    return sorted(out, key=lambda d: (d.get("address", ""), d.get("radius_m") or 0))


def build_workbook(reports: list, xlsx_path: Optional[Path] = None) -> Path:
    """Write the workbook for these report dicts. Links to each HTML report are
    relative, so keep the workbook in the same folder as the reports."""
    if not reports:
        raise ValueError("no zoning reports to export")
    xlsx_path = Path(xlsx_path or REPORT_DIR / WORKBOOK_NAME)
    xlsx_path.parent.mkdir(parents=True, exist_ok=True)
    book = _Book()
    wb = book.wb
    summary = wb.active
    summary.title = "Summary"
    used = {"summary", "applications", "permits", DEV_SHEET.lower(), "sources"}
    briefs = {lot_label(d): _sheet_name(lot_label(d), used) for d in reports}
    brief_ws = {lab: wb.create_sheet(name) for lab, name in briefs.items()}
    apps_ws = wb.create_sheet("Applications")
    permits_ws = wb.create_sheet("Permits")
    rez_ws = wb.create_sheet(DEV_SHEET)
    sources_ws = wb.create_sheet("Sources")

    with _body_pt(11):
        _write_summary(book, summary, reports, briefs)
        for d in reports:
            _write_brief(book, brief_ws[lot_label(d)], d, lot_label(d))
    outcome_style = {"Outcome": {v[0]: (v[1], v[2]) for v in OUTCOME.values()}}
    _record_sheet(apps_ws, "Committee of Adjustment applications near each lot",
                  "Every minor variance and consent file within each lot's radius. Filter Lot to see "
                  "one lot. Documents says where to get the file's papers (its AIC page, or CoA staff "
                  "and the Research Request Portal); Raw City data opens the Open Data row.", APP_COLS,
                  _application_rows(reports), "Applications", outcome_style)
    _record_sheet(permits_ws, "Building-permit projects near each lot",
                  "Permits grouped into projects by permit-number root. Documents opens the permit "
                  "status search while the City still lists the permit, else a building records "
                  "request; Raw City data opens the Open Data rows.", PERMIT_COLS, _permit_rows(reports), "Permits", {})
    _record_sheet(rez_ws, "Development applications near each lot",
                  "Rezonings and Official Plan amendments (OZ), site plans (SA), subdivisions (SB), "
                  "condominiums (CD) and part lot control (PL). Documents opens the application in "
                  "the Application Information Centre, or emails Development Review for closed ones; "
                  "Raw City data opens the Open Data row.", REZ_COLS, _rezoning_rows(reports),
                  "DevelopmentApplications", {})
    _write_sources(sources_ws, reports)

    wb.calculation.fullCalcOnLoad = True
    wb.properties.title = "Parcis Zoning Reports"
    wb.properties.creator = "Parcis Zoning Agent"
    wb.properties.description = NOTICE
    wb.properties.keywords = WATERMARK
    wb.active = 0
    wb.save(xlsx_path)
    sheet_files = {f"xl/worksheets/sheet{i}.xml": ws.title for i, ws in enumerate(wb.worksheets, start=1)}
    _inject_cached_values(xlsx_path, book.cache, sheet_files)
    return xlsx_path


def export_workbook(report_dir: Optional[Path] = None, xlsx_path: Optional[Path] = None,
                    extra: Optional[list] = None) -> Path:
    """Rebuild the workbook from every report JSON in `report_dir`, plus any `extra`
    report dicts not saved there (they replace a saved report with the same slug)."""
    report_dir = Path(report_dir or REPORT_DIR)
    reports = load_reports(report_dir)
    if extra:
        slugs = {d.get("slug") for d in extra}
        reports = [d for d in reports if d.get("slug") not in slugs] + list(extra)
        reports.sort(key=lambda d: (d.get("address", ""), d.get("radius_m") or 0))
    return build_workbook(reports, xlsx_path or report_dir / WORKBOOK_NAME)


def main(argv: Optional[list] = None) -> int:
    ap = argparse.ArgumentParser(description="Rebuild the zoning workbook from the saved report JSONs")
    ap.add_argument("--dir", type=Path, default=REPORT_DIR, help="folder with the <slug>.json reports")
    ap.add_argument("--out", type=Path, default=None, help=f"workbook path (default <dir>/{WORKBOOK_NAME})")
    args = ap.parse_args(argv)
    try:
        path = export_workbook(args.dir, args.out)
    except ValueError as e:
        print(f"No reports to export in {args.dir}: {e}", file=sys.stderr)
        return 1
    except PermissionError:
        print("Couldn't write the workbook. Close it in Excel and run this again.", file=sys.stderr)
        return 1
    print(path)
    return 0


if __name__ == "__main__":
    sys.exit(main())
