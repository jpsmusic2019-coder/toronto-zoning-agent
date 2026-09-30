# Copyright © 2026 Joshua Seaton. All rights reserved. Evaluation only; see LICENSE.
"""Zoning Agent tests — pure logic, outage handling, outputs. No network, no real DB."""
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path

import pytest

from toronto_zoning_agent import agent as za
from toronto_zoning_agent import gis
from toronto_zoning_agent import report as zh
from toronto_zoning_agent import rules as zr
from toronto_zoning_agent import data as zd

FIXTURE = json.loads((Path(__file__).parent / "data" / "ardmore_coa_250m.json").read_text(encoding="utf-8"))
RUN_DATE = date.fromisoformat(FIXTURE["run_date"])


def _coa(d: dict) -> zd.CoARecord:
    keys = {k: v for k, v in d.items() if k in zd.CoARecord.__dataclass_fields__}
    rec = zd.CoARecord(**keys)
    rec.address_norm = zd.addr_key(rec.street_num, rec.street_name, rec.street_type)
    return rec


NEIGHBOURS = [_coa(d) for d in FIXTURE["neighbours"]]


# ── EST_CONST_COST parser + bands ─────────────────────────────────────────────

@pytest.mark.unit
def test_est_cost_valid_in_band():
    value, unreliable = zd.parse_est_cost("1,000,000", "SFD - Detached")
    assert value == 1_000_000
    assert unreliable is False


@pytest.mark.unit
def test_est_cost_sentinel_is_null_and_unreliable():
    value, unreliable = zd.parse_est_cost("DO NOT UPDATE OR DELETE THIS INFO FIELD", "SFD - Detached")
    assert value is None
    assert unreliable is True


@pytest.mark.unit
def test_est_cost_zero_flagged_by_floor():
    value, unreliable = zd.parse_est_cost("0", "SFD - Detached")
    assert value == 0
    assert unreliable is True  # <= $1,000 floor


@pytest.mark.unit
def test_est_cost_over_ceiling_flagged_but_kept():
    value, unreliable = zd.parse_est_cost("30000000", "SFD - Detached")
    assert value == 30_000_000
    assert unreliable is True


@pytest.mark.unit
def test_est_cost_megaproject_passes_in_its_own_band():
    value, unreliable = zd.parse_est_cost("120000000", "Apartment Building")
    assert value == 120_000_000
    assert unreliable is False


@pytest.mark.unit
@pytest.mark.parametrize("structure,expected", [
    ("SFD - Detached", "single_family"),
    ("Laneway / Rear Yard Suite", "single_family"),
    ("3+ Unit - Townhouse", "multi_unit"),
    ("Apartment Building", "multi_unit"),
    ("Retail Store", "retail"),
    ("Office", "commercial"),
    ("Mixed Use/Res w Non Res", "mixed_use"),
    ("Warehouse", "industrial"),
    ("Parking Garage", "institutional_other"),
    ("", "institutional_other"),
])
def test_structure_category(structure, expected):
    assert zd.structure_category(structure) == expected


# ── permit <-> CoA decision-id linker + grouping ─────────────────────────────

@pytest.mark.unit
def test_extract_linked_coa_finds_decision_id():
    desc = "...See also 22 111400 DEM, 21 148009 MV and Final and Binding A0554/21TEY."
    assert zd.extract_linked_coa(desc) == ["A0554/21TEY"]


@pytest.mark.unit
def test_extract_linked_coa_ignores_mv_folder_only():
    assert zd.extract_linked_coa("Conforms to 21 148009 MV application.") == []


@pytest.mark.unit
def test_extract_linked_coa_dedupes_multiple():
    desc = "A0927/22TEY supersedes A0927/22TEY; see A1344/21TEY."
    assert zd.extract_linked_coa(desc) == ["A0927/22TEY", "A1344/21TEY"]


@pytest.mark.unit
def test_permit_root_groups_siblings():
    assert zd.permit_root("22 111410 BLD") == "22 111410"
    assert zd.permit_root("22 111410 HVA") == "22 111410"


@pytest.mark.unit
def test_is_new_build_permit_excludes_demolition():
    demo = zd.PermitRecord(permit_type="Demolition", structure_type="SFD - Detached")
    new_house = zd.PermitRecord(permit_type="New Houses", structure_type="SFD - Detached")
    laneway = zd.PermitRecord(permit_type="Small Residential Projects",
                              structure_type="Laneway / Rear Yard Suite")
    assert zd.is_new_build_permit(demo) is False
    assert zd.is_new_build_permit(new_house) is True
    assert zd.is_new_build_permit(laneway) is True


@pytest.mark.unit
@pytest.mark.parametrize("types,works,expected", [
    (["New Houses", "Plumbing(PS)"], ["New Building"], "New building"),
    (["Demolition Folder (DM)"], ["Demolition"], "Demolition"),
    (["Small Residential Projects"], ["Deck"], "Addition or alteration"),
    (["Building Additions/Alterations"], ["Interior Alterations"], "Addition or alteration"),
    (["Small Residential Projects"], ["New Building"], "New building"),
    (["Drain and Site Service"], ["Back Water Valve (Sewer only)"], "Other"),
])
def test_permit_category(types, works, expected):
    assert zr.permit_category(types, works) == expected


# ── address parsing ───────────────────────────────────────────────────────────

@pytest.mark.unit
@pytest.mark.parametrize("address,expected", [
    ("32 Ardmore Rd", ("32", "ARDMORE", "RD", "")),
    ("32 Ardmore Road", ("32", "ARDMORE", "ROAD", "")),
    ("32 Ardmore Rd, Toronto, ON M5P 1V6", ("32", "ARDMORE", "RD", "")),
    ("65 Bristol Ave", ("65", "BRISTOL", "AVE", "")),
    ("100 Queen St W", ("100", "QUEEN", "ST", "W")),
    ("100 Queen Street West", ("100", "QUEEN", "STREET", "W")),
    ("12A Glenayr Road", ("12A", "GLENAYR", "ROAD", "")),
    ("1203-5 Soudan Ave", ("5", "SOUDAN", "AVE", "")),
    ("Unit 1203, 5 Soudan Ave", ("5", "SOUDAN", "AVE", "")),
    ("5 Soudan Ave #1203", ("5", "SOUDAN", "AVE", "")),
    ("32 Ardmore Rd.", ("32", "ARDMORE", "RD", "")),
])
def test_parse_address(address, expected):
    assert za.parse_address(address) == expected


@pytest.mark.unit
def test_parse_address_input_keeps_unit_and_postal():
    p = za.parse_address_input("1203-5 Soudan Ave, Toronto, ON M4S 1V5")
    assert p["unit"] == "1203"
    assert p["fsa_hint"] == "M4S"


@pytest.mark.unit
def test_normalize_address():
    assert gis.normalize_address("  32  Ardmore   Rd ") == "32 ARDMORE RD"
    assert zd.addr_key("32", "Ardmore", "Rd", "None") == "32 ARDMORE RD"


_ARDMORE_ROWS = [{"key": "32 ARDMORE RD", "num": "32", "name": "ARDMORE", "type": "RD",
                  "dir": "", "full": "32 Ardmore Rd", "lat": 43.6948, "lon": -79.4184}]


@pytest.mark.unit
@pytest.mark.parametrize("text", ["32 Ardmore Rd", "32 Ardmore Road",
                                  "32 Ardmore Rd, Toronto, ON M5P 1V6"])
def test_resolve_subject_variants(monkeypatch, text):
    monkeypatch.setattr(gis, "find_address", lambda num, name: _ARDMORE_ROWS if name == "ARDMORE" else [])
    s = za.resolve_subject(text)
    assert s.key == "32 ARDMORE RD"
    assert s.full == "32 Ardmore Rd"


@pytest.mark.unit
def test_resolve_subject_unit_form(monkeypatch):
    rows = [{"key": "5 SOUDAN AVE", "num": "5", "name": "SOUDAN", "type": "AVE", "dir": "",
             "full": "5 Soudan Ave", "lat": 43.70, "lon": -79.39}]
    monkeypatch.setattr(gis, "find_address", lambda num, name: rows if (num, name) == ("5", "SOUDAN") else [])
    s = za.resolve_subject("1203-5 Soudan Ave")
    assert s.key == "5 SOUDAN AVE" and s.unit == "1203"


@pytest.mark.unit
def test_resolve_subject_not_in_toronto(monkeypatch):
    monkeypatch.setattr(gis, "find_address", lambda num, name: [])
    with pytest.raises(za.AddressNotFound):
        za.resolve_subject("1 Nowhere Lane")


@pytest.mark.unit
def test_cli_unknown_address_exits_cleanly(monkeypatch, capsys):
    monkeypatch.setattr(gis, "find_address", lambda num, name: [])
    monkeypatch.setattr(gis, "announce_downloads", lambda names: None)
    code = za.main(["--address", "1 Nowhere Lane", "--no-claude"])
    err = capsys.readouterr().err
    assert code == 2
    assert "Address not found" in err and "32 Ardmore Rd" in err
    assert "Traceback" not in err


# ── haversine ─────────────────────────────────────────────────────────────────

@pytest.mark.unit
def test_haversine_one_degree_latitude_is_about_111km():
    assert za._haversine_m(43.0, -79.41, 44.0, -79.41) == pytest.approx(111_195, rel=0.01)


# ── project type classifier (real 32 Ardmore descriptions) ────────────────────

@pytest.mark.unit
def test_project_types_match_32_ardmore():
    types = Counter(zr.project_type(c.application_type, c.sub_type, c.description) for c in NEIGHBOURS)
    assert len(NEIGHBOURS) == 45
    assert types == {"New house": 26, "Addition or alteration": 10,
                     "Pool, deck or exterior": 5, "Legalize existing work": 2,
                     "Severance or consent": 2}


@pytest.mark.unit
@pytest.mark.parametrize("ref,expected", [
    ("A0276/24TEY", "Pool, deck or exterior"),      # cabana + pool, Add/Alt sub-type
    ("A0473/26TEY", "Pool, deck or exterior"),      # concrete pad and walkway
    ("A0198/24TEY", "Pool, deck or exterior"),      # rear second storey balcony
    ("A0251/20TEY", "Legalize existing work"),      # landscaping, decks, shed
    ("A0026/22TEY", "Legalize existing work"),
    ("A0286/17TEY", "Addition or alteration"),      # third floor addition
    ("A0879/20TEY", "Addition or alteration"),      # "Part1", Add/Alt sub-type
    ("B0050/23TEY", "Severance or consent"),
    ("A0793/18TEY", "New house"),                   # New Res, "front basement extension"
])
def test_project_type_real_examples(ref, expected):
    c = next(c for c in NEIGHBOURS if c.reference_file == ref)
    assert zr.project_type(c.application_type, c.sub_type, c.description) == expected


@pytest.mark.unit
def test_storey_word_in_description_is_not_a_height_tag():
    # The old tagger called this "height"; it is a pool.
    assert zr.project_type("MV", "Add/Alt to Existing Res <= 3 units",
                           "To construct an outdoor swimming pool in the rear yard of the "
                           "existing three-storey detached house.") == "Pool, deck or exterior"


# ── storeys ───────────────────────────────────────────────────────────────────

@pytest.mark.unit
@pytest.mark.parametrize("text,expected", [
    ("To construct a new 2-½ storey detached dwelling", 2.5),
    ("a new 2 ½-storey house", 2.5),
    ("a new 2½ storey house", 2.5),
    ("a new detached two-and-a-half storey dwelling", 2.5),
    ("the existing two-and-one-half-storey dwelling", 2.5),
    ("a new three-storey detached dwelling", 3.0),
    ("a new 3 storey dwelling", 3.0),
    ("a new two-storey detached dwelling", 2.0),
    ("a one-storey ancillary structure", 1.0),
    ("a new detached dwelling", None),
])
def test_parse_storeys(text, expected):
    assert zr.parse_storeys(text) == expected


@pytest.mark.unit
def test_approved_new_house_storeys_32_ardmore():
    approved_nh = [c for c in NEIGHBOURS
                   if zr.outcome(c.decision, c.source_resource)[0] == "Approved"
                   and zr.project_type(c.application_type, c.sub_type, c.description) == "New house"]
    assert len(approved_nh) == 22
    assert Counter(zr.parse_storeys(c.description) for c in approved_nh) == {3.0: 15, 2.5: 2, 2.0: 5}


# ── outcomes, decided counts, signal ─────────────────────────────────────────

@pytest.mark.unit
@pytest.mark.parametrize("decision,source,expected", [
    ("Approved", "closed", ("Approved", "Approved")),
    ("Approved with Conditions", "active", ("Approved", "Approved with Conditions")),
    ("Refused", "closed", ("Refused", "Refused")),
    ("Deferred", "closed", ("Other", "Deferred")),
    ("Withdrawn", "closed", ("Other", "Withdrawn")),
    ("", "active", ("Pending", "Awaiting hearing")),
    ("", "closed", ("Other", "Closed, no decision recorded")),
])
def test_outcome_labels(decision, source, expected):
    assert zr.outcome(decision, source) == expected


@pytest.mark.unit
def test_decided_counts_32_ardmore():
    buckets = Counter(zr.outcome(c.decision, c.source_resource)[0] for c in NEIGHBOURS)
    decided = buckets["Approved"] + buckets["Refused"]
    assert (buckets["Approved"], decided) == (37, 39)
    assert zr.precedent_signal(buckets["Approved"], decided) == "Strong"


@pytest.mark.unit
@pytest.mark.parametrize("approved,decided,expected", [
    (2, 2, "Sparse"),     # fewer than 3 decided
    (0, 0, "Sparse"),
    (8, 10, "Strong"),    # 80%
    (37, 39, "Strong"),
    (5, 10, "Mixed"),     # 50%
    (7, 10, "Mixed"),     # 70%
    (4, 10, "Weak"),      # 40%
    (0, 3, "Weak"),
])
def test_precedent_signal(approved, decided, expected):
    assert za.precedent_signal(approved, decided) == expected


# ── time to hearing ───────────────────────────────────────────────────────────

@pytest.mark.unit
def test_hearing_timing_32_ardmore():
    rows = []
    for c in NEIGHBOURS:
        b, _ = zr.outcome(c.decision, c.source_resource)
        if zr.is_decided(b):
            rows.append((c.in_date, zr.days_between(c.in_date, c.hearing_date)))
    t = zr.hearing_timing(rows, RUN_DATE)
    assert t["cutoff"] == "2023-09-28"
    assert t["recent_n"] == 9
    assert zr.weeks(t["recent_median_days"]) == 8
    assert zr.weeks(t["all_median_days"]) == 16
    assert zr.typical_weeks(t) == (8, True)


@pytest.mark.unit
def test_hearing_timing_falls_back_to_all_years():
    t = zr.hearing_timing([("2018-01-01", 70), ("2018-02-01", 84), ("2019-01-01", 98),
                           ("2026-01-01", 30)], date(2026, 9, 28))
    assert t["recent_n"] == 1
    assert zr.typical_weeks(t) == (11, False)


@pytest.mark.unit
def test_days_between():
    assert zr.days_between("2021-04-30T00:00:00", "2021-09-29") == 152
    assert zr.days_between("2021-04-30", "") is None


# ── zone decoder ──────────────────────────────────────────────────────────────

@pytest.mark.unit
def test_decode_zone_rd():
    z = za.decode_zone("RD (f12.0; d0.65) (x1321) (ZZC)")
    assert z["raw"] == "RD (f12.0; d0.65) (x1321)"
    assert z["name"] == "Residential Detached"
    assert (z["frontage_m"], z["fsi"], z["exception"]) == (12.0, 0.65, 1321)


@pytest.mark.unit
def test_decode_zone_area_and_units():
    z = zr.decode_zone("RD (f15.0; a550) (x5)")
    assert z["min_lot_area_m2"] == 550.0
    z = zr.decode_zone("RM (u4; d0.8)")
    assert z["max_units"] == 4 and z["fsi"] == 0.8


@pytest.mark.unit
def test_decode_zone_cr():
    z = zr.decode_zone("CR 3.0 (c2.0; r2.5) SS2 (x1234)")
    assert z["name"] == "Commercial Residential"
    assert (z["total_fsi"], z["commercial_fsi"], z["residential_fsi"]) == (3.0, 2.0, 2.5)
    assert (z["standard_set"], z["exception"]) == (2, 1234)


@pytest.mark.unit
def test_decode_zone_fallback_to_raw():
    z = zr.decode_zone("see by-law 438-86")
    assert z["parsed"] is False and z["raw"] == "see by-law 438-86"


# ── subject chain wording ─────────────────────────────────────────────────────

@pytest.mark.unit
def test_subject_chain_target_wording():
    subj = _coa(FIXTURE["subject"]["coa"])
    permit = zd.PermitRecord(permit_num="22 111410 BLD", permit_root="22 111410",
                             permit_type="New Houses", structure_type="SFD - Detached",
                             work="New Building", application_date="2022-02-07",
                             issued_date="2022-02-23", status="Inspection",
                             linked_coa=["A0554/21TEY"])
    _, chain = za.build_subject_chain([subj], za._group_permits([permit]))
    assert chain == ("Minor variance A0554/21TEY approved 29 Sep 2021, 152 days after filing; "
                     "new-house permit 22 111410 issued 23 Feb 2022, now under inspection.")


@pytest.mark.unit
def test_refused_then_approved_chain():
    refused = zd.CoARecord(reference_file="A1344/21TEY", application_type="MV", decision="Refused",
                           in_date="2021-10-01", hearing_date="2022-02-23", source_resource="closed")
    approved = zd.CoARecord(reference_file="A0927/22TEY", application_type="MV", decision="Approved",
                            in_date="2022-07-06", hearing_date="2022-11-09", source_resource="closed")
    _, chain = za.build_subject_chain([refused, approved], [])
    assert chain.startswith("Minor variance A1344/21TEY was refused 23 Feb 2022; "
                            "a revised application, A0927/22TEY, was approved 9 Nov 2022")


@pytest.mark.unit
def test_no_history_subject_chain_is_empty():
    assert za.build_subject_chain([], []) == ("", "")


# ── outage handling (build_report with the City sources faked) ────────────────

def _fake_gis(monkeypatch, subject_lat=43.6948, subject_lon=-79.4184):
    monkeypatch.setattr(za, "resolve_subject", lambda a: za.Subject(
        input=a, num="32", name="ARDMORE", stype="RD", direction="", unit="",
        key="32 ARDMORE RD", full="32 Ardmore Rd", lat=subject_lat, lon=subject_lon))
    coords = {c.address_norm: (c.lat, c.lon) for c in NEIGHBOURS}
    coords["32 ARDMORE RD"] = (subject_lat, subject_lon)
    monkeypatch.setattr(gis, "lookup_addresses", lambda keys: {k: coords[k] for k in keys if k in coords})
    zone = {"zoning_area": {"zone": "RD", "zone_string": "RD (f12.0; d0.65) (x1321)", "gen_zone": "0"},
            "height": {"max_height_m": "12", "height_string": "HT 12.0"}}
    monkeypatch.setattr(gis, "point_lookup", lambda layer, lat, lon: zone.get(layer))
    monkeypatch.setattr(gis, "nearest_distance_m", lambda layer, lat, lon, search_m=5000: 1636.0)
    monkeypatch.setattr(gis, "features_in_box", lambda layer, bbox, **k: [])
    monkeypatch.setattr(gis, "nearby_streets", lambda lat, lon, radius_m=200, limit=6: [])
    monkeypatch.setattr(zd, "aic_index", lambda rsns: {})  # no network in tests


def _build():
    return za.build_report("32 Ardmore Rd", 250, "test", True, True, True, False,
                           district="C03", run_date=RUN_DATE)


_NO_PRECEDENT_WORDS = re.compile(r"untested|no precedent|first-mover|none found|no committee", re.I)


@pytest.mark.unit
def test_all_sources_down_reads_unchecked(monkeypatch):
    _fake_gis(monkeypatch)
    monkeypatch.setenv("ZONING_CKAN_OFFLINE", "1")
    r = _build()
    m = r.memo
    for k in ("subject_coa", "subject_permits", "neighbour_coa", "neighbour_permits", "dev_apps"):
        assert r.status[k] == "unchecked"
    assert "couldn't be checked" in m.summary
    assert not _NO_PRECEDENT_WORDS.search(m.summary)
    assert "neighbour_coa:unchecked" in m.data_completeness
    assert m.zone_code == "RD (f12.0; d0.65) (x1321)"  # local GIS still works


@pytest.mark.unit
def test_fsa_resolves_then_coa_pull_fails(monkeypatch):
    _fake_gis(monkeypatch)
    monkeypatch.setattr(zd, "resolve_fsa", lambda name, num="": "M5P")

    def boom(fsa):
        raise zd.CKANUnavailable("CKAN unreachable for coa_closed: timeout")
    monkeypatch.setattr(zd, "coa_records_for_fsa", boom)
    monkeypatch.setattr(zd, "permit_records_for_fsa", lambda fsa: [])
    monkeypatch.setattr(zd, "devapp_records_for_fsa", lambda fsa: [])
    r = _build()  # must not raise
    assert r.status["neighbour_coa"] == "unchecked"
    assert r.status["subject_coa"] == "unchecked"
    assert r.status["neighbour_permits"] == "none"
    assert "coa" in r.memo.partial_flags
    assert "Neighbour precedent couldn't be checked" in r.memo.summary
    assert "untested" not in r.memo.summary


@pytest.mark.unit
def test_real_empty_fsa_reads_none_found(monkeypatch):
    _fake_gis(monkeypatch)
    monkeypatch.setattr(zd, "resolve_fsa", lambda name, num="": "M5P")
    monkeypatch.setattr(zd, "coa_records_for_fsa", lambda fsa: [])
    monkeypatch.setattr(zd, "permit_records_for_fsa", lambda fsa: [])
    monkeypatch.setattr(zd, "devapp_records_for_fsa", lambda fsa: [])
    r = _build()
    assert r.status["neighbour_coa"] == "none"
    assert r.memo.unchecked_sources == []
    assert "No Committee of Adjustment applications are on record within 250 m" in r.memo.summary
    assert "couldn't be checked" not in r.memo.summary


@pytest.mark.unit
def test_resolve_fsa_raises_when_ckan_unreachable(monkeypatch):
    monkeypatch.setenv("ZONING_CKAN_OFFLINE", "1")
    with pytest.raises(zd.CKANUnavailable):
        zd.resolve_fsa("ARDMORE", "32")


@pytest.mark.unit
def test_resolve_fsa_none_when_ckan_answers_empty(monkeypatch):
    monkeypatch.setattr(zd, "fetch_datastore", lambda *a, **k: {"records": [], "total": 0})
    assert zd.resolve_fsa("NOWHERE", "1") is None


def _fixture_report(monkeypatch):
    _fake_gis(monkeypatch)
    subj = _coa(FIXTURE["subject"]["coa"])
    monkeypatch.setattr(zd, "resolve_fsa", lambda name, num="": "M5P")
    monkeypatch.setattr(zd, "coa_records_for_fsa", lambda fsa: [subj] + [_coa(d) for d in FIXTURE["neighbours"]])
    monkeypatch.setattr(zd, "permit_records_for_fsa", lambda fsa: [])
    monkeypatch.setattr(zd, "devapp_records_for_fsa", lambda fsa: [])
    return _build()


@pytest.mark.unit
def test_build_report_32_ardmore_numbers(monkeypatch):
    r = _fixture_report(monkeypatch)
    m = r.memo
    assert (m.neighbour_count, m.approved_nearby, m.decided_nearby) == (45, 37, 39)
    assert m.precedent_signal == "Strong"
    assert (m.recent_hearing_weeks, m.all_hearing_weeks) == (8, 16)
    assert m.district == "C03"  # --district reaches the memo
    assert "Within 250 m, 37 of 39 decided applications were approved (95%), including 15 " \
           "three-storey new houses." in m.summary
    assert "dominant" not in m.summary.lower()


# ── outputs: JSON contract, template render, Excel summary ───────────────────

REQUIRED_KEYS = {"schema_version", "generated", "address", "fsa", "lat", "lon", "radius_m",
                 "links", "zone", "envelope", "summary", "stats", "subject", "records",
                 "permits", "devapps", "completeness", "unchecked_sources", "map", "sources",
                 "licence"}


@pytest.mark.unit
def test_report_dict_required_keys(monkeypatch):
    d = zh.report_to_dict(_fixture_report(monkeypatch), RUN_DATE)
    assert REQUIRED_KEYS <= set(d)
    assert len(d["records"]) == 45
    rec = d["records"][0]
    assert {"id", "ref", "addr", "d", "lat", "lon", "type", "storeys", "outcome", "label",
            "filed", "hearing", "days", "desc", "permits", "links"} <= set(rec)
    assert d["licence"].startswith("Contains information licensed under the Open Government Licence")
    assert d["schema_version"] == "1.3"
    assert rec["docs"][0]["action"] == "coa_staff" and rec["docs"][0]["number"] == rec["ref"]
    assert rec["dataset"].startswith("Committee of Adjustment Applications")
    assert rec["src"].startswith("https://ckan0.cf.opendata.inter.prod-toronto.ca/api/3/action/datastore_search?")
    assert {x["key"] for x in d["datasets"]} >= {"coa", "permits_active", "permits_cleared", "dev_apps"}
    assert d["zone"]["links"]["chapter"]["url"].endswith("ZBL_NewProvision_Chapter10_20.htm")
    assert d["zone"]["links"]["exception"]["url"].endswith("ZBL_NewProvision_Chapter900_3.htm")
    json.dumps(d)  # serializable


@pytest.mark.unit
def test_template_render_json_block_parses(monkeypatch):
    d = zh.report_to_dict(_fixture_report(monkeypatch), RUN_DATE)
    d["records"][0]["desc"] += " </script><script>alert(1)</script> <!-- x"
    page = zh.render_html(d)
    m = re.search(r'<script type="application/json" id="report-data">(.*?)</script>', page, re.S)
    assert m, "report-data block missing"
    assert json.loads(m.group(1)) == d
    assert "__REPORT_DATA__" not in page and "/*__LEAFLET_JS__*/" not in page
    assert "Leaflet 1.9.4" in page
    assert "cdn.jsdelivr" not in page and "unpkg.com" not in page
    assert "fonts.googleapis" not in page and "/*__FONTS_CSS__*/" not in page
    assert page.count("@font-face") == len(zh.FONTS)


def _report_json(monkeypatch, tmp_path, address="32 Ardmore Rd"):
    """A saved report JSON (as the agent writes it) for the fixture lot."""
    r = _fixture_report(monkeypatch)
    r.memo.address = address
    d = zh.report_to_dict(r, RUN_DATE, workbook="zoning_reports.xlsx")
    (tmp_path / f"{d['slug']}.json").write_text(json.dumps(d), encoding="utf-8")
    return d


def _all_cells(ws):
    return [c for row in ws.iter_rows() for c in row]


@pytest.mark.unit
def test_workbook_lives_beside_the_reports_and_lists_every_lot(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    _report_json(monkeypatch, tmp_path)
    _report_json(monkeypatch, tmp_path, "65 Bristol Ave")
    path = zx.export_workbook(tmp_path)
    assert path == tmp_path / "zoning_reports.xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "32 Ardmore Rd", "65 Bristol Ave", "Applications",
                             "Permits", "Development Applications", "Sources"]
    assert [c.value for c in wb["Summary"]["A"][5:7]] == ["32 Ardmore Rd", "65 Bristol Ave"]
    # rebuilt from the JSONs on disk: dropping a report drops its rows and sheet
    (tmp_path / "65-bristol-ave.json").unlink()
    wb = load_workbook(zx.export_workbook(tmp_path))
    assert "65 Bristol Ave" not in wb.sheetnames
    assert {c.value for c in wb["Applications"]["A"][4:]} == {"32 Ardmore Rd"}


@pytest.mark.unit
def test_workbook_links_are_real_hyperlinks(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    d = _report_json(monkeypatch, tmp_path)
    wb = load_workbook(zx.export_workbook(tmp_path))
    formulas = [c.value for ws in wb.worksheets for c in _all_cells(ws)
                if isinstance(c.value, str) and c.value.upper().startswith("=HYPERLINK")]
    assert formulas == []  # no =HYPERLINK(): those need recalculation and break in viewers
    s = wb["Summary"]
    assert s["A6"].hyperlink.location == "'32 Ardmore Rd'!A1"
    report = next(c for c in _all_cells(s) if c.value == "Open report")
    assert report.hyperlink.target == "32-ardmore-rd.html"  # relative: works from any clone
    apps = wb["Applications"]
    head = [c.value for c in apps[4]]
    raw_col, city_col = head.index("Raw City data"), head.index("Documents")
    assert raw_col == len(head) - 1  # raw data is the last column
    assert city_col == head.index("File") + 1  # the human link sits beside the file number
    rows = list(apps.iter_rows(min_row=5))
    assert len(rows) == len(d["records"]) == 45  # every neighbour, no 60-row cap
    for row in rows:
        link = row[raw_col].hyperlink
        assert link and link.target.startswith(
            "https://ckan0.cf.opendata.inter.prod-toronto.ca/api/3/action/datastore_search?")
        assert row[city_col].hyperlink.target.startswith(("https://www.toronto.ca/", "https://secure.toronto.ca/", "mailto:"))
        assert row[city_col].value in ("Open on City site", "Request from CoA staff")
    assert "REFERENCE_FILE%23" in rows[0][raw_col].hyperlink.target


@pytest.mark.unit
def test_workbook_counts_are_formulas_with_cached_values(monkeypatch, tmp_path):
    import zipfile
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    _report_json(monkeypatch, tmp_path)
    path = zx.export_workbook(tmp_path)
    f = load_workbook(path)["Summary"]
    v = load_workbook(path, data_only=True)["Summary"]
    head = {c.value: c.column_letter for c in f[5]}
    assert f[f"{head['Approved']}6"].value.startswith("=COUNTIFS(Applications!")
    assert v[f"{head['Approved']}6"].value == 37
    assert v[f"{head['Decided']}6"].value == 39
    assert v[f"{head['Approval rate']}6"].value == pytest.approx(37 / 39)
    assert v[f"{head['3-storey new houses approved']}6"].value == 15
    assert f[f"{head['Approval rate']}6"].number_format == "0%"
    with zipfile.ZipFile(path) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert not any("<v />" in z.read(n).decode() or "<v/>" in z.read(n).decode() for n in sheets)


@pytest.mark.unit
def test_workbook_types_dates_and_wording(monkeypatch, tmp_path):
    from datetime import date as _date
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    _report_json(monkeypatch, tmp_path)
    wb = load_workbook(zx.export_workbook(tmp_path))
    apps = wb["Applications"]
    head = [c.value for c in apps[4]]
    first = [c.value for c in apps[5]]
    assert isinstance(first[head.index("Dist. (m)")], int)
    assert isinstance(first[head.index("Filed")], _date)  # a real date, not text
    outcomes = {c.value for c in apps[apps[4][head.index("Outcome")].column_letter][4:]}
    assert outcomes <= {"Approved", "Refused", "Awaiting hearing", "Not decided"}
    labels = [c.value for c in apps[apps[4][head.index("Decision on file")].column_letter][4:]]
    assert labels.count("Closed, no decision recorded") == 3
    assert "Awaiting hearing" in labels
    brief = [str(c.value) for c in _all_cells(wb["32 Ardmore Rd"]) if c.value is not None]
    assert any("No lot coverage overlay at this lot" in v for v in brief)
    assert any(v.startswith("Strong precedent: 37 of 39 decided") for v in brief)
    assert not any("District" == v for v in (c.value for c in wb["Summary"][5]))


@pytest.mark.unit
def test_workbook_marks_unchecked_sources_not_zero(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    d = _report_json(monkeypatch, tmp_path)
    d["completeness"]["neighbour_coa"] = "unchecked"
    d["unchecked_sources"] = ["neighbour_coa"]
    d["stats"]["signal"] = "Unchecked"
    (tmp_path / "32-ardmore-rd.json").write_text(json.dumps(d), encoding="utf-8")
    s = load_workbook(zx.export_workbook(tmp_path))["Summary"]
    head = {c.value: c.column_letter for c in s[5]}
    assert s[f"{head['Approved']}6"].value == "n/a ⚠"
    assert s[f"{head['Precedent']}6"].value == "Couldn't be checked"
    assert "couldn't be checked" in s[f"{head['Data check']}6"].value


@pytest.mark.unit
def test_persist_keeps_every_neighbour(monkeypatch, tmp_path):
    import sqlite3
    r = _fixture_report(monkeypatch)
    db = tmp_path / "t.db"
    za.persist(r, db)
    za.persist(r, db)  # rerun replaces rows
    n = sqlite3.connect(db).execute(
        "SELECT COUNT(*) FROM zoning_coa WHERE is_subject = 0").fetchone()[0]
    assert n == 45
    cols = {row[1] for row in sqlite3.connect(db).execute("PRAGMA table_info(zoning_coa)")}
    assert "project_type" in cols


@pytest.mark.unit
def test_persist_migrates_old_schema(tmp_path, monkeypatch):
    import sqlite3
    db = tmp_path / "old.db"
    conn = sqlite3.connect(db)
    conn.execute("CREATE TABLE zoning_coa (id INTEGER PRIMARY KEY AUTOINCREMENT, client_id TEXT, "
                 "subject_address TEXT, reference_file TEXT, variance_tag TEXT, "
                 "source_resource TEXT, street_num TEXT, "
                 "UNIQUE(client_id, subject_address, reference_file, source_resource, street_num))")
    conn.commit()
    conn.close()
    za.persist(_fixture_report(monkeypatch), db)
    cols = {row[1] for row in sqlite3.connect(db).execute("PRAGMA table_info(zoning_coa)")}
    assert {"project_type", "storeys", "variance_tag"} <= cols


@pytest.mark.unit
def test_outage_run_keeps_previous_rows(monkeypatch, tmp_path):
    import sqlite3
    db = tmp_path / "t.db"
    za.persist(_fixture_report(monkeypatch), db)

    def down_fsa(name, num=""):
        raise zd.CKANUnavailable("offline")
    monkeypatch.setattr(zd, "resolve_fsa", down_fsa)
    down = _build()
    assert down.status["neighbour_coa"] == "unchecked"
    za.persist(down, db)
    conn = sqlite3.connect(db)
    assert conn.execute("SELECT COUNT(*) FROM zoning_coa WHERE is_subject = 0").fetchone()[0] == 45
    assert "neighbour_coa:unchecked" in conn.execute(
        "SELECT data_completeness FROM zoning_memos").fetchone()[0]


@pytest.mark.unit
def test_fsa_pulls_use_exact_postal_filter(monkeypatch):
    calls = []
    monkeypatch.setattr(zd, "fetch_all", lambda rid, q=None, filters=None, cap=60000: calls.append((q, filters)) or [])
    zd.coa_records_for_fsa("M5H")
    zd.permit_records_for_fsa("M5H")
    zd.devapp_records_for_fsa("M5H")
    assert calls and all(q is None and f == {"POSTAL": "M5H"} for q, f in calls)


@pytest.mark.unit
def test_claude_prompt_carries_facts_and_unchecked_sources(monkeypatch):
    from toronto_zoning_agent import narrative as zn
    sent = {}

    class FakeMessages:
        def create(self, **kw):
            sent.update(kw)
            return type("R", (), {"stop_reason": "end_turn",
                                  "content": [type("B", (), {"type": "text", "text": " Interpretation. "})()]})()

    fake = type("C", (), {"messages": FakeMessages()})()
    monkeypatch.setattr(zn, "_get_client", lambda: fake)
    facts = {"address": "32 Ardmore Rd", "radius_m": 250, "neighbour_status": "found",
             "neighbour_coa_count": 45, "decided_nearby": 39, "approved_nearby": 37,
             "approval_pct": 95, "three_storey_new_houses": 15, "approved_new_houses": 22,
             "typical_hearing_weeks": 8, "typical_hearing_recent": True,
             "unchecked_sources": ["neighbour building permits"]}
    text, version = zn.generate_zoning_narrative(facts, use_claude=True)
    prompt = sent["messages"][0]["content"]
    assert text == "Interpretation." and version.startswith("zoning-narrative-claude")
    assert "37 of 39 decided applications were approved (95%)" in prompt
    assert "could NOT be checked this run: neighbour building permits" in prompt
    assert "Do NOT invent" in prompt
    assert "temperature" not in sent and sent["extra_body"] == {"temperature": 0.3}


@pytest.mark.unit
def test_claude_failure_reason_is_reported(monkeypatch):
    from toronto_zoning_agent import narrative as zn

    class Boom(Exception):
        status_code = 401
        body = {"type": "error", "error": {"type": "authentication_error", "message": "invalid x-api-key"}}

    class FakeMessages:
        def create(self, **kw):
            raise Boom("Error code: 401")

    monkeypatch.setattr(zn, "_get_client", lambda: type("C", (), {"messages": FakeMessages()})())
    text, version = zn.generate_zoning_narrative({"radius_m": 250}, use_claude=True)
    assert version == zn.TEMPLATE_VERSION
    assert zn.last_failure == "401 invalid x-api-key"
    zn.generate_zoning_narrative({"radius_m": 250}, use_claude=False)
    assert zn.last_failure == ""


@pytest.mark.unit
def test_build_report_flags_claude_failure(monkeypatch, capsys):
    from toronto_zoning_agent import narrative as zn
    _fake_gis(monkeypatch)
    monkeypatch.setattr(zd, "resolve_fsa", lambda name, num="": "M5P")
    for f in ("coa_records_for_fsa", "permit_records_for_fsa", "devapp_records_for_fsa"):
        monkeypatch.setattr(zd, f, lambda fsa: [])
    monkeypatch.setattr(zn, "_get_client", lambda: None)
    r = za.build_report("32 Ardmore Rd", 250, "test", True, True, True, True, run_date=RUN_DATE)
    assert r.memo.partial_flags["narrative"] == "no ANTHROPIC_API_KEY set"
    out = capsys.readouterr().out
    assert "No ANTHROPIC_API_KEY set, so the summary uses the template" in out
    assert "failed" not in out and r.memo.summary_source == "template"


# ── source links + narrative fact check ──────────────────────────────────────

@pytest.mark.unit
def test_source_links_point_at_the_right_city_records():
    from urllib.parse import parse_qs, urlparse
    u = urlparse(zd.coa_record_url("A0554/21TEY", "closed"))
    q = parse_qs(u.query)
    assert q["resource_id"] == [zd.RESOURCES["coa_closed"]]
    assert json.loads(q["filters"][0]) == {"REFERENCE_FILE#": "A0554/21TEY"}
    q = parse_qs(urlparse(zd.coa_record_url("A0568/26TEY", "active")).query)
    assert q["resource_id"] == [zd.RESOURCES["coa_active"]]
    q = parse_qs(urlparse(zd.permit_records_url(["22 111410 PLB", "22 111410 BLD"], "active")).query)
    assert json.loads(q["filters"][0]) == {"PERMIT_NUM": ["22 111410 BLD", "22 111410 PLB"]}
    assert zd.coa_record_url("", "closed") == ""
    q = parse_qs(urlparse(zd.devapp_record_url("25 240145 STE 09 OZ")).query)
    assert q["resource_id"] == [zd.RESOURCES["dev_apps"]]
    assert json.loads(q["filters"][0]) == {"APPLICATION#": "25 240145 STE 09 OZ"}
    assert zd.bylaw_links("R", 739)["exception"]["label"] == "Exception R 739, Chapter 900.2"
    assert zd.bylaw_links("XYZ", None)["chapter"]["url"] == zd.BYLAW_URL


@pytest.mark.unit
def test_permit_projects_keep_their_numbers_for_links():
    groups = za._group_permits([
        zd.PermitRecord(permit_num="22 111410 BLD", permit_root="22 111410", permit_class="active"),
        zd.PermitRecord(permit_num="22 111410 PLB", permit_root="22 111410", permit_class="cleared"),
        zd.PermitRecord(permit_num="22 111410 DRN", permit_root="22 111410", permit_class="active"),
    ])
    rec = zh._permit_record({**groups[0], "address": "32 Ardmore Rd"}, "p0")
    assert rec["nums"] == ["22 111410 BLD", "22 111410 DRN", "22 111410 PLB"]
    assert zd.RESOURCES["permits_active"] in rec["src"]
    assert zd.RESOURCES["permits_cleared"] in rec["src_other"]


@pytest.mark.unit
def test_narrative_timing_is_called_a_median():
    from toronto_zoning_agent import narrative as zn
    facts = {"typical_hearing_weeks": 8, "typical_hearing_recent": True, "recent_hearing_n": 9}
    assert "median" in zn.timing_sentence(facts)
    assert "median 8 weeks" in zn._claude_prompt(facts)
    assert "never call it\nan average" in zn._claude_prompt(facts)


@pytest.mark.unit
def test_narrative_check_rejects_averages_and_invented_numbers():
    from toronto_zoning_agent import narrative as zn
    prompt = "37 of 39 decided (95%) within 250 m; median 8 weeks; A0554/21TEY 29 Sep 2021"
    assert zn.check_narrative("37 of 39 approved (95%) within 250 m, about 8 weeks.", prompt) == []
    assert zn.check_narrative("An 8-week average hearing time.", prompt) == [
        "called the median hearing time an average"]
    assert zn.check_narrative("Roughly 40 approvals.", prompt) == ["numbers not in the facts: 40"]


@pytest.mark.unit
def test_narrative_retries_once_then_falls_back(monkeypatch):
    from toronto_zoning_agent import narrative as zn
    replies = iter(["Hearings average 8 weeks.", "Hearings take a median 8 weeks."])
    sent = []

    class FakeMessages:
        def create(self, **kw):
            sent.append(kw["messages"][0]["content"])
            text = next(replies)
            return type("R", (), {"stop_reason": "end_turn",
                                  "content": [type("B", (), {"type": "text", "text": text})()]})()

    monkeypatch.setattr(zn, "_get_client", lambda: type("C", (), {"messages": FakeMessages()})())
    facts = {"radius_m": 250, "typical_hearing_weeks": 8, "typical_hearing_recent": True}
    text, version = zn.generate_zoning_narrative(facts, use_claude=True)
    assert text == "Hearings take a median 8 weeks." and version.startswith("zoning-narrative-claude")
    assert len(sent) == 2 and "previous draft was rejected" in sent[1]

    replies = iter(["Hearings average 8 weeks.", "Still an average of 8 weeks."])
    text, version = zn.generate_zoning_narrative(facts, use_claude=True)
    assert version == zn.TEMPLATE_VERSION
    assert zn.last_failure.startswith("RuntimeError: draft failed the fact check")


@pytest.mark.unit
def test_rounding_matches_the_page():
    # The page uses Math.round (half up); Python's round() is half-to-even.
    assert zr.round_half_up(12.5) == 13
    assert zr.round_half_up(10.4999) == 10
    assert zr.weeks(73.5) == 11
    assert zr.weeks(70) == 10


# ── links to the City's own pages ─────────────────────────────────────────────

@pytest.mark.unit
def test_document_routes_coa_in_and_out_of_the_aic():
    run = date(2026, 9, 29)
    aic = {"5888249": "290129"}
    inside = zd.document_routes("coa", "A0568/26TEY", run, "5888249", aic, "210 Rosemary Rd",
                                decided_on="2026-08-01")
    assert [r["action"] for r in inside] == ["aic"]
    assert inside[0]["url"].endswith("?id=5888249&pid=290129&title=210-ROSEMARY-RD")
    recent = zd.document_routes("coa", "A0385/18TEY", run, "123", aic, decided_on="2018-09-26")
    assert [r["action"] for r in recent] == ["coa_staff", "research"]
    assert "$150 + HST" in recent[1]["detail"] and "$300 + HST" in recent[1]["detail"]
    assert recent[0]["url"] == zd.COA_STAFF_URL and recent[1]["url"] == zd.RESEARCH_REQUEST_URL


@pytest.mark.unit
def test_document_routes_coa_older_than_ten_years_is_staff_only():
    run = date(2026, 9, 29)
    old = zd.document_routes("coa", "A0100/15TEY", run, decided_on="2016-09-28")
    assert [r["action"] for r in old] == ["coa_staff"] and "older than 10 years" in old[0]["detail"]
    edge = zd.document_routes("coa", "A0100/16TEY", run, decided_on="2016-09-29")
    assert [r["action"] for r in edge] == ["coa_staff", "research"]  # exactly 10 years: still in


@pytest.mark.unit
def test_permit_status_window_edges():
    run = date(2026, 9, 29)
    w = zd.permit_in_status_window
    assert w("Inspection", "2016-09-29", "", run)          # open, exactly 10 years
    assert not w("Inspection", "2016-09-28", "", run)      # open, a day past 10 years
    assert w("Closed", "2014-03-01", "2026-08-29", run)    # closed exactly a month ago
    assert not w("Closed", "2014-03-01", "2026-08-28", run)
    assert not w("Closed", "2014-03-01", "", run)          # closed, no date: not listed
    assert w("Cancelled", "2025-01-01", "2026-09-20", run)
    assert not w("Closed - Dormant", "2023-01-01", "2024-01-01", run)
    old = zd.document_routes("permit", "14 135787", run, permits=[("Closed", "2014-05-01", "2015-02-01")])
    assert [r["action"] for r in old] == ["records"]
    assert old[0]["email"] == "bldrecords@toronto.ca" and "$76.98" in old[0]["detail"]
    live = zd.document_routes("permit", "22 111410", run,
                              permits=[("Closed", "2022-02-07", "2023-01-01"), ("Inspection", "2022-02-07", "")])
    assert [r["action"] for r in live] == ["status"] and live[0]["url"] == zd.PERMIT_STATUS_URL


@pytest.mark.unit
def test_document_routes_devapps():
    run = date(2026, 9, 29)
    aic = {"5725239": "173426"}
    dev = zd.document_routes("devapp", "25 240145 STE 09 OZ", run, "5725239", aic, "1423 Dufferin St")
    assert dev[0]["action"] == "aic"
    closed = zd.document_routes("devapp", "14 213788 WET 17 SA", run, "999", aic)
    assert closed[0]["action"] == "devreview" and closed[0]["url"].startswith("mailto:developmentreview@toronto.ca")
    assert zd.devapp_type_name("SA") == "Site Plan Approval"
    assert zd.devapp_type_name("OZ") == "Official Plan Amendment / Rezoning"
    assert zd.devapp_is_open("NOAC Issued") and not zd.devapp_is_open("Closed")


@pytest.mark.unit
def test_published_fields_drop_contacts_and_read_plainly():
    rec = zd._coa_from_raw({"REFERENCE_FILE#": "A0554/21TEY", "IN_DATE": "2021-04-30T00:00:00",
                            "CONTACT_NAME": "x", "CONTACT_EMAIL": "x@y", "CONTACT_PHONE": "1",
                            "NUMBER_OF_LOTS_CREATED": "2", "APPLICATION_URL": "http://app.toronto.ca/AIC",
                            "APPLICATION_TYPE": "MV", "SYS_ID": "4917671"}, "closed")
    assert not {"CONTACT_NAME", "CONTACT_EMAIL", "CONTACT_PHONE", "APPLICATION_URL"} & set(rec.published)
    f = dict(zh.plain_fields(rec.published, zh.COA_FIELDS))
    assert f == {"File number": "A0554/21TEY", "Application type": "Minor variance",
                 "Filed": "30 Apr 2021", "Lots created": "2"}


@pytest.mark.unit
def test_aic_index_offline_raises_unchecked_not_empty(monkeypatch):
    monkeypatch.setenv("ZONING_CKAN_OFFLINE", "1")
    assert zd.aic_index([]) == {}
    with pytest.raises(zd.CKANUnavailable):
        zd.aic_index(["4917671"])


@pytest.mark.unit
def test_no_primary_link_points_at_the_ckan_api(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from toronto_zoning_agent import export as zx
    d = _report_json(monkeypatch, tmp_path)
    groups = [d["records"], d["permits"], d["devapps"], d["subject"]["coa"], d["subject"]["permits"]]
    links = [r for g in groups for x in g for r in x["docs"]]
    assert links and all(link["url"] and "/api/3/action/" not in link["url"] for link in links)
    assert "/api/3/action/" not in d["links"]["zoning_map"]
    wb = load_workbook(zx.export_workbook(tmp_path))
    for ws in wb.worksheets:
        head = [c.value for c in ws[4]] if ws.title in ("Applications", "Permits", "Development Applications") else []
        raw = head.index("Raw City data") + 1 if "Raw City data" in head else None
        for c in _all_cells(ws):
            t = c.hyperlink.target if c.hyperlink else None
            if t and "/api/3/action/" in t:
                assert raw and c.column == raw and c.value == "Raw data", (ws.title, c.coordinate, c.value)


@pytest.mark.unit
def test_zoning_map_link_lands_on_the_lot():
    url = zd.zoning_map_url(43.694811, -79.418365)
    assert url == "https://map.toronto.ca/gccmaps/?app=zoning&center=-79.418365,43.694811&scale=1128"
    assert zd.zoning_map_url(None, None) == zd.ZONING_MAP_URL


@pytest.mark.unit
def test_narrative_check_rejects_a_restated_status():
    from toronto_zoning_agent import narrative as zn
    prompt = "This lot's history: permit 22 111410 issued, with a revision issued."
    assert zn.check_narrative("The permit is now under revision.", prompt) == [
        'restated a status as "under revision"; use the status word for word']
    assert zn.check_narrative("The permit was issued, with a revision issued.", prompt) == []
    assert "word for word" in zn._claude_prompt({"radius_m": 250})
    assert zn.check_narrative("**Developer attention:** the permit was issued.", prompt) == [
        "used markdown formatting; write plain sentences"]
    assert zn.check_narrative("2 three-storey new houses in the same R (d0.6) zone.", "2 three-storey 0.6") == [
        "said the neighbours share the lot's zone, which the facts don't give"]
    assert zn.check_narrative("new houses within the R (d0.6) (x739) zone.", "0.6 739") != []
    assert zn.check_narrative("The lot is zoned RD (f12.0; d0.65).", "12.0 0.65") == []
    assert zn.check_narrative("Applications will be evaluated against the base zone framework.",
                              "Zone: RD (f12.0; d0.65) (x1321)") == [
        "said base zone rules apply, but the lot has a site-specific exception"]
    bad = "It may modify the floor space index parameters (f12.0; d0.65)."
    assert zn.check_narrative(bad, "12.0 0.65") == [
        "called the frontage token (f…) a floor space index; f is minimum lot frontage"]
    assert zn.check_narrative("A 12.0 m minimum frontage (f12.0) and a 0.65 FSI (d0.65).", "12.0 0.65") == []
    assert "f12.0 = Minimum lot frontage 12.0 m" in zn._claude_prompt(
        {"zone_decoded": "f12.0 = Minimum lot frontage 12.0 m; d0.65 = Maximum floor space index 0.65"})


@pytest.mark.unit
def test_default_tiles_are_the_city_basemap(monkeypatch):
    for k in ("ZONING_TILE_URL", "ZONING_TILE_URL_DARK", "ZONING_TILE_ATTRIBUTION"):
        monkeypatch.delenv(k, raising=False)
    t = zh.tile_config()
    assert t["url"].startswith("https://gis.toronto.ca/arcgis/rest/services/basemap/cot_topo/")
    assert "{z}/{y}/{x}" in t["url"] and "openstreetmap" not in t["url"]
    assert "City of Toronto" in t["attribution"]
    monkeypatch.setenv("ZONING_TILE_URL", "https://tiles.example/{z}/{x}/{y}.png")
    assert zh.tile_config()["url"] == "https://tiles.example/{z}/{x}/{y}.png"


@pytest.mark.unit
def test_template_counts_only_oz_as_rezonings():
    from toronto_zoning_agent import narrative as zn
    facts = {"radius_m": 250, "devapp_count": 3, "oz_count": 1, "devapp_status": "found"}
    text = zn.template_narrative(facts) if hasattr(zn, "template_narrative") else zn.generate_zoning_narrative(facts, False)[0]
    assert "3 development applications (1 rezoning or Official Plan amendment) within 250 m." in text
    assert "3 rezoning" not in text


@pytest.mark.unit
def test_resulting_storeys_only_when_stated():
    add = "Addition or alteration"
    assert zr.resulting_storeys(add, "To alter the approved building permit plans for alterations to a "
                                     "two-storey detached dwelling by constructing a third floor addition.") == 3
    assert zr.resulting_storeys(add, "To alter the existing one-storey detached dwelling by constructing a "
                                     "second storey addition.") == 2
    assert zr.resulting_storeys(add, "To construct a second-storey addition over the garage.") is None
    assert zr.resulting_storeys(add, "To construct a second storey rear addition.") is None
    assert zr.resulting_storeys(add, "To alter the ancillary building (detached garage) by constructing a "
                                     "complete second storey addition.") is None
    assert zr.resulting_storeys(add, "To construct a 3rd storey addition.") == 3
    assert zr.resulting_storeys(add, "To construct a rear two-storey addition to the dwelling.") is None
    assert zr.resulting_storeys(add, "To alter the existing two-storey detached dwelling by reconstructing the "
                                     "front bay window, and by constructing a third storey, rear ground floor addition.") == 3
    assert zr.resulting_storeys(add, "To alter the dwelling by constructing a rear two-storey addition with "
                                     "a rear second storey deck.") is None
    assert zr.resulting_storeys(add, "To construct a second storey deck and a rear balcony.") is None
    assert zr.resulting_storeys(add, "To alter the existing two-storey detached dwelling.") is None
    assert zr.resulting_storeys("New house", "To construct a new three-storey detached dwelling.") == 3
    assert zr.resulting_storeys("Pool, deck or exterior", "A one-storey cabana and a third floor addition.") is None
